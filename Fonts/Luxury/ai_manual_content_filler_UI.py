import sys
import subprocess
import time

if "streamlit" not in sys.modules:
    import os

    script = os.path.abspath(__file__)
    subprocess.Popen([sys.executable, "-m", "streamlit", "run", script])
    time.sleep(2)
    sys.exit()

import os
import re
import openpyxl
import streamlit as st

# ─── НАСТРОЙКИ ────────────────────────────────────────────────────────────────
START_ROW = 2  # С какой строки начнать запись
END_ROW = 5001  # До какой строки включительно писать
COL_DESC = 12  # L столбец для Описания
COL_CTA = 13  # M столбец для CTA
COL_TITLE = 14  # N столбец для Названия

# Доступные сферы
SPHERES = [
    "BEAUTY", "FASHION", "HOME", "TRAVEL", "FOOD_AND_DRINK",
    "DIY_AND_CRAFT", "HEALTH_AND_FITNESS", "EDUCATION",
    "DESIGN_AND_ART", "EVENTS", "OTHER"
]

st.set_page_config(
    page_title="🪄 Pinterest Content Filler",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=JetBrains+Mono:wght@400;700&family=Syne:wght@700;800&display=swap');

html, body, [class*="css"] {
    background-color: #0d0d14 !important;
    color: #e0e0f0 !important;
    font-family: 'JetBrains Mono', monospace !important;
}
section[data-testid="stSidebar"] {
    background: #111119 !important;
    border-right: 1px solid #1e1e2e !important;
}
section[data-testid="stSidebar"] * { color: #e0e0f0 !important; }
.main-header {
    font-family: 'Syne', sans-serif;
    font-size: 26px; font-weight: 800;
    background: linear-gradient(90deg, #a78bfa, #f472b6);
    -webkit-background-clip: text;
    -webkit-text-fill-color: transparent;
    margin-bottom: 2px;
}
.sub-header { color: #555580; font-size: 11px; margin-bottom: 16px; }
.card {
    background: #13131e; border: 1px solid #1e1e30;
    border-radius: 10px; padding: 16px 20px; margin-bottom: 12px;
}
.card-accent { border-left: 3px solid #a78bfa; }
.sec-label {
    font-size: 10px; font-weight: 700; letter-spacing: 1.5px;
    text-transform: uppercase; margin-bottom: 6px;
}
.col-purple { color: #a78bfa; }
.col-pink   { color: #f472b6; }
.prompt-box {
    background: #0a0a12; border: 1px solid #252538; border-radius: 8px;
    padding: 14px 18px; font-family: 'JetBrains Mono', monospace;
    font-size: 12px; line-height: 1.7; color: #c0c0e0;
    white-space: pre-wrap; word-break: break-word;
}
.niche-badge {
    display: inline-block; background: #1a1a2e;
    border: 1px solid #2e2e4e; border-radius: 5px;
    padding: 2px 10px; font-size: 11px; color: #a78bfa; margin-bottom: 5px;
}
.strat-title { font-family: 'Syne', sans-serif; font-size: 20px; font-weight: 700; color: #e0e0f0; }
.xlsx-path { font-size: 10px; color: #34d399; margin-top: 3px; }
div[data-testid="stTextArea"] textarea {
    background: #0a0a12 !important; color: #c0c0e0 !important;
    border: 1px solid #252538 !important; border-radius: 8px !important;
    font-family: 'JetBrains Mono', monospace !important; font-size: 12px !important;
}
div[data-testid="stTextArea"] textarea:focus {
    border-color: #a78bfa !important;
    box-shadow: 0 0 0 2px #a78bfa33 !important;
}
div.stButton > button {
    background: #1a1a2e !important; color: #e0e0f0 !important;
    border: 1px solid #2e2e4e !important; border-radius: 8px !important;
    font-family: 'JetBrains Mono', monospace !important; font-size: 12px !important;
}
div.stButton > button:hover {
    background: #252545 !important; border-color: #a78bfa !important; color: #c4b5fd !important;
}
div.stButton > button[kind="primary"] {
    background: linear-gradient(135deg, #7c3aed, #a855f7) !important;
    border: none !important; color: white !important; font-weight: 700 !important;
}
[data-testid="metric-container"] {
    background: #13131e; border: 1px solid #1e1e30;
    border-radius: 8px; padding: 10px 14px;
}
hr { border-color: #1e1e30 !important; }
#MainMenu, footer { visibility: hidden; }
header[data-testid="stHeader"] { background: #0d0d14 !important; }
</style>
""", unsafe_allow_html=True)


# ─── HELPERS ──────────────────────────────────────────────────────────────────
def read_txt(path):
    if not os.path.exists(path):
        return ""
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        return f.read().strip()


def clean_lines(raw):
    lines = []
    for line in raw.splitlines():
        line = line.strip()
        line = re.sub(r'^\d+[\.\)]\s*|^[*\-•]\s*|^`{3,}', '', line)
        line = line.replace("```", "").strip()
        if line:
            lines.append(line)
    return lines


def scan_tasks(base_dir):
    tasks = []
    if not os.path.isdir(base_dir):
        return tasks
    niches = sorted(
        d for d in os.listdir(base_dir)
        if os.path.isdir(os.path.join(base_dir, d))
        and not d.startswith(".")
        and d not in {"__pycache__", ".venv", "venv", ".git"}
    )
    for niche in niches:
        np = os.path.join(base_dir, niche)
        try:
            strats = sorted(
                d for d in os.listdir(np)
                if os.path.isdir(os.path.join(np, d)) and re.match(r"^\d+", d)
            )
        except PermissionError:
            continue
        for s in strats:
            sp = os.path.join(np, s)
            xlsx = next(
                (os.path.join(sp, f) for f in os.listdir(sp)
                 if f.endswith(".xlsx") and not f.startswith("~$")), None
            )
            tasks.append({"niche": niche, "strat": s, "path": sp, "xlsx": xlsx})
    return tasks


def check_xlsx_filled(xlsx_path):
    """Проверяет заполнены ли колонки L/M/N начиная со START_ROW."""
    if not xlsx_path:
        return False
    try:
        wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
        ws = wb.active
        filled = 0
        for row in range(START_ROW, min(START_ROW + 3, END_ROW + 1)):
            if ws.cell(row=row, column=COL_TITLE).value and \
                    ws.cell(row=row, column=COL_DESC).value and \
                    ws.cell(row=row, column=COL_CTA).value:
                filled += 1
        wb.close()
        return filled >= 3
    except Exception:
        return False


def read_unique_xlsx_data(xlsx_path):
    """Надежно считывает уникальные строки, чтобы не выводить дубли."""
    if not xlsx_path or not os.path.exists(xlsx_path):
        return "", "", ""

    titles, descs, ctas = [], [], []
    seen_t, seen_d, seen_c = set(), set(), set()

    try:
        # data_only=True - читаем значения, без read_only для надежности
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
        ws = wb.active

        # Сканируем только первые 500 строк, чтобы найти уники
        # (нет смысла читать все 5000, так как они дублируются)
        max_search_row = min(START_ROW + 500, ws.max_row + 1)

        for row in range(START_ROW, max_search_row):
            d_val = ws.cell(row=row, column=COL_DESC).value
            c_val = ws.cell(row=row, column=COL_CTA).value
            t_val = ws.cell(row=row, column=COL_TITLE).value

            # Собираем уникальные Заголовки
            if t_val is not None:
                cln = str(t_val).strip()
                if cln and cln not in seen_t:
                    titles.append(cln)
                    seen_t.add(cln)

            # Собираем уникальные Описания
            if d_val is not None:
                cln = str(d_val).strip()
                if cln and cln not in seen_d:
                    descs.append(cln)
                    seen_d.add(cln)

            # Собираем уникальные CTA
            if c_val is not None:
                cln = str(c_val).strip()
                if cln and cln not in seen_c:
                    ctas.append(cln)
                    seen_c.add(cln)

        wb.close()
        return "\n".join(titles), "\n".join(descs), "\n".join(ctas)
    except Exception as e:
        print(f"Ошибка чтения файла: {e}")
        return "", "", ""


def build_prompt(strat_path, base_dir):
    for_main = os.path.join(strat_path, "For Main")
    link = read_txt(os.path.join(base_dir, "linksite.txt"))
    tags = read_txt(os.path.join(for_main, "Boards_Tags.txt"))
    board = read_txt(os.path.join(for_main, "Boards_Name.txt"))
    missing = []
    if not link:  missing.append("linksite.txt")
    if not tags:  missing.append("Boards_Tags.txt")
    if not board: missing.append("Boards_Name.txt")
    if missing:
        return "", missing
    prompt = (
        f"Привет. Я продвигаю на Pinterest товар.\n"
        f"Ссылка на сайт (детально изучи данный сайт, чтобы понять что пользователь получит когда перейдет на данный сайт с фото Pinterest): {link}\n"
        f"Теги по которым продвигаюсь (используй их как контекст для ключевых слов, не вставляй теги дословно в тексты):\n"
        f"{tags}\n"
        f"Главный тег доски: {board}\n\n"
        f"Сделай на английском:\n"
        f"1) 100 ЗАГОЛОВКОВ (без эмодзи, без знаков препинания в конце, 3-8 слов, содержат ключевые слова из тегов естественно вписанные)\n"
        f"2) 100 ОПИСАНИЙ (с разными эмодзи по смыслу к содержанию и тегам в начале предложения, без артикля \"A\" в начале, без призыва к действию внутри, 1-2 предложения, должны подходить к ЛЮБОЙ картинке из выдачи по данным тегам)\n"
        f"3) 100 ПРИЗЫВОВ К ДЕЙСТВИЮ — CTA (с разными эмодзи по смыслу, одно короткое предложение, призыв перейти скачать наш товар бесплатно на сайте с подборкой лучшего товара, разные формулировки — не повторяться)\n\n"
        f"ВАЖНЫЕ ПРАВИЛА:\n"
        f"- Любой заголовок + любое описание + любой CTA должны сочетаться друг с другом\n"
        f"- Все тексты в рамках тематики тегов и товара, без выхода за рамки\n"
        f"- Каждый список помести в отдельный блок кода\n"
        f"- Без нумерации, без пустых строк между строками\n"
        f"- В CTA добаявляй с заглавной буквы акцент на выгодность перехода на наш сайт, то есть помечай что они получать при переходе на наш сайт по имени товара"
        f"- Так же в CTA пиши в 70% случаев FREE (с акцентом на этом и большими буквами), так они могут получить сейчас уже бесплатно то что им нужно - сделай акцент на FREE с большой буквы и чтобы видно было\n"
        f"- (ПРАВИЛО КАСАТЕЛЬНО CTA блок про - ПРОГРАММНОГО ОБЕСПЕЧЕНИЯ/ИНСТРУМЕНТОВ — КРИТИЧЕСКИ ВАЖНО): Если в тегах упоминается ПО (Canva, Adobe, Procreate, Cricut, GoodNotes и т. д.), НЕ пишите фразы определнных фирм программно обеспечения в CTA вроде «Fonts for Canva» или «Typography Adobe». "
        f" Вместо этого делайте акцент на преимуществах, исходя из темы нашей подборки тегов и товаров по ссылке выше. Используйте такие формулировки, которые *идеально подходят* для этих инструментов. Найдите «золотую середину» между тегами и самим продуктом."
        f"- Только сами тексты, ничего лишнего\n"
        f"- Так же сделай здесь чтобы не было подвязки под определенный вид тега (потому что нужно чтобы рандомное фото с рандомного тега взяли и оно подошло под эту нишу)"
    )
    return prompt, []


def write_xlsx(xlsx_path, titles, descs, ctas):
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active
    total = END_ROW - START_ROW + 1
    for i in range(total):
        row = START_ROW + i
        ws.cell(row=row, column=COL_DESC).value = descs[i % len(descs)]
        ws.cell(row=row, column=COL_CTA).value = ctas[i % len(ctas)]
        ws.cell(row=row, column=COL_TITLE).value = titles[i % len(titles)]
    wb.save(xlsx_path)


def count_badge(n):
    if n == 0:   return "<span style='color:#555580'>0 строк</span>"
    if n >= 90:  return f"<span style='color:#34d399'>✅ {n} строк</span>"
    if n >= 50:  return f"<span style='color:#fbbf24'>⚠️ {n} строк</span>"
    return f"<span style='color:#f87171'>❌ {n} строк</span>"


def update_text_areas(task_idx):
    """Обновляет состояние текстовых полей значениями из Excel."""
    if not st.session_state.tasks: return
    task = st.session_state.tasks[task_idx]
    t, d, c = read_unique_xlsx_data(task["xlsx"])
    st.session_state.ta_titles = t
    st.session_state.ta_descs = d
    st.session_state.ta_ctas = c


# ─── СОСТОЯНИЕ ────────────────────────────────────────────────────────────────
base_dir = os.path.dirname(os.path.abspath(__file__))

if "tasks" not in st.session_state:
    st.session_state.tasks = scan_tasks(base_dir)
if "done_set" not in st.session_state:
    done = set()
    for i, t in enumerate(st.session_state.tasks):
        if check_xlsx_filled(t["xlsx"]):
            done.add(i)
    st.session_state.done_set = done

if "cur_idx" not in st.session_state:
    all_idx = set(range(len(st.session_state.tasks)))
    remaining = sorted(all_idx - st.session_state.done_set)
    st.session_state.cur_idx = remaining[0] if remaining else 0

if "last_idx" not in st.session_state:
    st.session_state.last_idx = st.session_state.cur_idx

# Подгружаем данные при первичном запуске
if "ta_titles" not in st.session_state:
    update_text_areas(st.session_state.cur_idx)

# Загружаем данные из таблицы при смене задачи
if st.session_state.last_idx != st.session_state.cur_idx:
    update_text_areas(st.session_state.cur_idx)
    st.session_state.last_idx = st.session_state.cur_idx

tasks = st.session_state.tasks
cur = st.session_state.cur_idx

# ─── SIDEBAR ──────────────────────────────────────────────────────────────────
with st.sidebar:
    st.markdown('<div class="main-header">🪄 Filler</div>', unsafe_allow_html=True)
    st.markdown('<div class="sub-header">Pinterest Content Pipeline</div>', unsafe_allow_html=True)

    if tasks:
        total = len(tasks)
        done_count = len(st.session_state.done_set)
        st.progress(done_count / total if total else 0,
                    text=f"{done_count} / {total} выполнено")

    st.divider()

    if not tasks:
        st.warning("Папки стратегий не найдены рядом со скриптом.")
    else:
        last_niche = None
        for i, t in enumerate(tasks):
            if t["niche"] != last_niche:
                st.markdown(f"**📁 {t['niche']}**")
                last_niche = t["niche"]

            is_done = i in st.session_state.done_set
            has_xlsx = bool(t["xlsx"])
            icon = "✅" if is_done else ("📊" if has_xlsx else "❌")

            if i == cur:
                st.markdown(
                    f"<div style='background:#1e1030;border:1px solid #a78bfa;"
                    f"border-radius:6px;padding:5px 10px;font-size:12px;color:#c4b5fd;"
                    f"margin-bottom:3px'>▶ {t['strat']}</div>",
                    unsafe_allow_html=True
                )
            else:
                if st.button(f"{icon}  {t['strat']}", key=f"nav_{i}", use_container_width=True):
                    st.session_state.cur_idx = i
                    st.rerun()

    st.divider()
    if st.button("🔄  Пересканировать папки", use_container_width=True):
        st.session_state.tasks = scan_tasks(base_dir)
        done = set()
        for i, t in enumerate(st.session_state.tasks):
            if check_xlsx_filled(t["xlsx"]):
                done.add(i)
        st.session_state.done_set = done
        all_idx = set(range(len(st.session_state.tasks)))
        remaining = sorted(all_idx - done)
        st.session_state.cur_idx = remaining[0] if remaining else 0
        update_text_areas(st.session_state.cur_idx)
        st.session_state.last_idx = st.session_state.cur_idx
        st.rerun()

# ─── MAIN ─────────────────────────────────────────────────────────────────────
if not tasks:
    st.markdown('<div class="main-header">🪄 Pinterest Content Filler</div>', unsafe_allow_html=True)
    st.error(f"Папки стратегий не найдены.\n\nПуть скрипта: `{base_dir}`")
    st.stop()

task = tasks[cur]

# Заголовок задачи
xlsx_text = f"📊 {task['xlsx']}" if task["xlsx"] else "❌ XLSX не найден"
st.markdown(f"""
<div class="card card-accent">
  <div class="niche-badge">📁 {task['niche']}</div>
  <div class="strat-title">{task['strat']}</div>
  <div class="xlsx-path">{xlsx_text}</div>
</div>
""", unsafe_allow_html=True)

# Навигация
c1, c2, c3 = st.columns([1, 1, 6])
with c1:
    if st.button("◀  Назад", disabled=(cur == 0), use_container_width=True):
        st.session_state.cur_idx -= 1
        st.rerun()
with c2:
    if st.button("Вперёд  ▶", disabled=(cur >= len(tasks) - 1), use_container_width=True):
        st.session_state.cur_idx += 1
        st.rerun()

st.divider()

# ─── БЛОК ВЫБОРА СФЕРЫ ────────────────────────────────────────────────────────
st.markdown('<div class="sec-label col-purple">📌 Настройка сферы (For Main/Sphere.txt)</div>', unsafe_allow_html=True)

for_main_dir = os.path.join(task["path"], "For Main")
sphere_path = os.path.join(for_main_dir, "Sphere.txt")
current_sphere = read_txt(sphere_path).strip()

# Подготавливаем опции и текущее значение
if not current_sphere:
    options = ["⚠️ НЕ ВЫБРАНО"] + SPHERES
    current_val = "⚠️ НЕ ВЫБРАНО"
elif current_sphere not in SPHERES:
    options = [current_sphere] + SPHERES
    current_val = current_sphere
else:
    options = SPHERES
    current_val = current_sphere

col_sph1, col_sph2 = st.columns([2, 2])
with col_sph1:
    selected_sphere = st.selectbox(
        "Выберите нишу (изменение сохраняется в файл автоматически):",
        options=options,
        index=options.index(current_val),
        key=f"sphere_select_{cur}"
    )

if selected_sphere != current_val and not selected_sphere.startswith("⚠️"):
    os.makedirs(for_main_dir, exist_ok=True)
    with open(sphere_path, "w", encoding="utf-8") as f:
        f.write(selected_sphere)
    st.toast(f"✅ Файл Sphere.txt обновлен: {selected_sphere}")
    st.rerun()

st.divider()

# ─── ПРОМПТ ───────────────────────────────────────────────────────────────────
st.markdown('<div class="sec-label col-purple">1️⃣  Промпт для AI — скопируй кнопкой ⎘ справа</div>',
            unsafe_allow_html=True)

prompt, missing = build_prompt(task["path"], base_dir)
if missing:
    st.error(f"⚠️ Не найдены файлы: `{'`, `'.join(missing)}`")
else:
    with st.expander("📋  Показать / скрыть промпт", expanded=False):
        st.code(prompt, language=None)

st.divider()

# ─── ПОЛЯ ВВОДА / ВЫВОДА ──────────────────────────────────────────────────────
head_col1, head_col2 = st.columns([4, 1])
with head_col1:
    st.markdown(
        '<div class="sec-label col-pink" style="margin-top:10px;">2️⃣  Вставь ответ из AI (или просмотри текущие данные)</div>',
        unsafe_allow_html=True)
with head_col2:
    # Кнопка ручной подгрузки, если вдруг авто-загрузка не сработала
    if st.button("📥 Прочитать из XLSX", use_container_width=True):
        with st.spinner("Считываю данные из таблицы..."):
            update_text_areas(cur)
        st.rerun()

col1, col2, col3 = st.columns(3)

with col1:
    raw_titles = st.text_area("📝 ЗАГОЛОВКИ", height=340, key="ta_titles",
                              placeholder="Вставьте заголовки из AI...\n\nНумерация убирается автоматически.")
    titles = clean_lines(raw_titles)
    st.markdown(count_badge(len(titles)), unsafe_allow_html=True)

with col2:
    raw_descs = st.text_area("📖 ОПИСАНИЯ", height=340, key="ta_descs",
                             placeholder="Вставьте описания из AI...")
    descs = clean_lines(raw_descs)
    st.markdown(count_badge(len(descs)), unsafe_allow_html=True)

with col3:
    raw_ctas = st.text_area("🎯 CTA", height=340, key="ta_ctas",
                            placeholder="Вставьте CTA из AI...")
    ctas = clean_lines(raw_ctas)
    st.markdown(count_badge(len(ctas)), unsafe_allow_html=True)

st.divider()

if st.session_state.get("_clear_fields"):
    st.session_state["ta_titles"] = ""
    st.session_state["ta_descs"] = ""
    st.session_state["ta_ctas"] = ""
    st.session_state["_clear_fields"] = False

# Метрики + запись
m1, m2, m3, _, btn_col = st.columns([1, 1, 1, 1, 2])
m1.metric("Заголовков", len(titles))
m2.metric("Описаний", len(descs))
m3.metric("CTA", len(ctas))

with btn_col:
    write_ready = bool(titles and descs and ctas and task["xlsx"])
    if st.button("💾  ЗАПИСАТЬ В XLSX", type="primary",
                 use_container_width=True, disabled=not write_ready):
        try:
            with st.spinner("Записываю данные..."):
                write_xlsx(task["xlsx"], titles, descs, ctas)
            st.session_state.done_set.add(cur)
            st.success(f"✅ Сохранено: `{os.path.basename(task['xlsx'])}`")
            st.balloons()
        except Exception as e:
            st.error(f"❌ Ошибка записи: {e}")

if not write_ready:
    if not task["xlsx"]:
        st.caption("❌ В папке стратегии нет XLSX файла")
    else:
        st.caption("⬆️ Заполни все три поля чтобы активировать кнопку записи")