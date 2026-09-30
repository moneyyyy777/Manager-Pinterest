import sys
import subprocess
import os

# 1. ПРОВЕРЯЕМ АВТОЗАПУСК СТРОГО ДО ИМПОРТА STREAMLIT!
if "streamlit" not in sys.modules:
    script = os.path.abspath(__file__)
    subprocess.run([sys.executable, "-m", "streamlit", "run", script])
    sys.exit()

# 2. ИМПОРТИРУЕМ ВСЁ ОСТАЛЬНОЕ
import json
import time
import re
import openpyxl
import streamlit as st
from playwright.sync_api import sync_playwright, TimeoutError

# ─── НАСТРОЙКИ ────────────────────────────────────────────────────────────────
START_ROW = 2
END_ROW = 5001
COL_DESC = 12
COL_CTA = 13
COL_TITLE = 14

CHROME_DEBUG_URL = "http://127.0.0.1:9222"

st.set_page_config(page_title="🤖 Pinterest Filler & QC", layout="wide", initial_sidebar_state="expanded")

st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=JetBrains+Mono:wght@400;700&family=Syne:wght@700;800&display=swap');
html, body, [class*="css"] { background-color: #0d0d14 !important; color: #e0e0f0 !important; font-family: 'JetBrains Mono', monospace !important; }
section[data-testid="stSidebar"] { background: #111119 !important; border-right: 1px solid #1e1e2e !important; }

/* ─── ТОЧЕЧНЫЙ БЕЛЫЙ ЦВЕТ ДЛЯ 3 ЭЛЕМЕНТОВ В САЙДБАРЕ ─── */
section[data-testid="stSidebar"] h3 { color: #ffffff !important; }
section[data-testid="stSidebar"] .stProgress p { color: #ffffff !important; }
section[data-testid="stSidebar"] [data-testid="stCaptionContainer"] p { color: #ffffff !important; }

.niche-badge {
    display: inline-block; 
    background: #25183e !important;
    border: 1px solid #a78bfa !important; 
    border-radius: 6px;
    padding: 6px 14px; 
    font-size: 13px; 
    font-weight: 700;
    color: #ffffff !important; 
    margin-bottom: 8px;
}
.main-header { font-family: 'Syne', sans-serif; font-size: 26px; font-weight: 800; background: linear-gradient(90deg, #a78bfa, #f472b6); -webkit-background-clip: text; -webkit-text-fill-color: transparent; }
.card { background: #13131e; border: 1px solid #1e1e30; border-radius: 10px; padding: 16px 20px; margin-bottom: 12px; }
.card-accent { border-left: 3px solid #a78bfa; }
div[data-testid="stTextArea"] textarea { background: #0a0a12 !important; color: #c0c0e0 !important; border: 1px solid #252538 !important; }
div.stButton > button[kind="primary"] { background: linear-gradient(135deg, #7c3aed, #a855f7) !important; border: none !important; color: white !important; font-weight: 700 !important; }
.stRadio > div { background: #1a1a2e; padding: 10px; border-radius: 8px; border: 1px solid #2e2e4e; margin-bottom: 10px; }
.stRadio * { color: #ffffff !important; }
</style>
""", unsafe_allow_html=True)


# ─── ИНТЕГРАЦИЯ С PLAYWRIGHT (AI STUDIO) ──────────────────────────────────────
def interact_with_ai_studio(prompt: str, expected_blocks: int, json_file_path: str, max_attempts: int = 15):
    with sync_playwright() as p:
        try:
            browser = p.chromium.connect_over_cdp(CHROME_DEBUG_URL)
            context = browser.contexts[0]

            page = None
            for p_ext in context.pages:
                if "aistudio" in p_ext.url:
                    page = p_ext
                    break
            if not page:
                page = context.new_page()

            for attempt in range(1, max_attempts + 1):
                try:
                    page.bring_to_front()
                    st.toast(f"🔄 Попытка {attempt}/{max_attempts}: Подготовка AI Studio...")

                    page.goto("https://aistudio.google.com/app/u/1/prompts/new_chat", timeout=30000)
                    page.wait_for_timeout(2000)

                    if json_file_path and os.path.exists(json_file_path):
                        try:
                            file_input = page.locator('input[type="file"]').first
                            file_input.set_input_files(json_file_path)
                            page.wait_for_timeout(3000)
                        except Exception as e:
                            print(f"⚠️ Ошибка загрузки файла: {e}")

                    text_area = page.locator("textarea").last
                    text_area.fill(prompt)
                    page.wait_for_timeout(500)

                    try:
                        grounding_sw = page.locator('button[role="switch"][aria-label*="Grounding with Google Search"]')
                        if grounding_sw.is_visible(timeout=500) and grounding_sw.get_attribute(
                                "aria-checked") == "true":
                            grounding_sw.click()
                            page.wait_for_timeout(300)

                        url_sw = page.locator('button[role="switch"][aria-label*="Browse the url context"]')
                        if url_sw.is_visible(timeout=500) and url_sw.get_attribute("aria-checked") == "true":
                            url_sw.click()
                            page.wait_for_timeout(300)
                    except Exception:
                        pass

                    try:
                        clicked = page.evaluate('''() => {
                            const btn = document.querySelector('button[mattooltipclass="run-button-tooltip"]') 
                                     || document.querySelector('button.ctrl-enter-submits');
                            if (btn) { btn.click(); return true; }
                            return false;
                        }''')
                        if not clicked:
                            page.locator('button[mattooltipclass="run-button-tooltip"]').first.click(force=True)
                    except Exception:
                        text_area.focus()
                        page.keyboard.press("Meta+Enter")

                    st.toast(f"⏳ Попытка {attempt}: AI генерирует ответ...")

                    prev_len = 0
                    stable_ticks = 0

                    for _ in range(90):  # До 3 минут генерации
                        time.sleep(2)

                        page_content = page.content()
                        if "An internal error has occurred" in page_content or "Failed to generate content" in page_content:
                            raise Exception("AI Studio выдал внутреннюю ошибку.")

                        code_blocks = page.locator("ms-code-block pre").all_inner_texts()
                        if not code_blocks:
                            code_blocks = page.locator("pre").all_inner_texts()

                        if len(code_blocks) >= expected_blocks:
                            curr_len = len(code_blocks[-1])
                            if curr_len > 30 and curr_len == prev_len:
                                stable_ticks += 1
                                if stable_ticks >= 3:
                                    final_blocks = page.locator("ms-code-block pre").all_inner_texts()
                                    if not final_blocks:
                                        final_blocks = page.locator("pre").all_inner_texts()
                                    return final_blocks[-expected_blocks:]
                            else:
                                stable_ticks = 0
                                prev_len = curr_len

                    raise Exception("Таймаут ожидания ответа AI.")

                except Exception as attempt_err:
                    st.warning(f"⚠️ Попытка {attempt} сбой: {attempt_err}. Перезапуск через 5 сек...")
                    time.sleep(5)

            raise Exception(f"Все {max_attempts} попыток завершились ошибкой AI Studio.")
        except Exception as e:
            st.error(f"❌ Ошибка Playwright: {e}")
            return []


# ─── ХЕЛПЕРЫ ──────────────────────────────────────────────────────────────────
def read_txt(path):
    if not os.path.exists(path): return ""
    with open(path, "r", encoding="utf-8", errors="replace") as f: return f.read().strip()


def clean_lines(raw):
    if not raw: return []
    lines = []
    for line in raw.splitlines():
        line = line.strip()
        line = re.sub(r'^\d+[\.\)]\s*|^[*\-•]\s*|^`{3,}', '', line).replace("```", "").strip()
        if line: lines.append(line)
    return lines


def scan_tasks(base_dir):
    tasks = []
    if not os.path.isdir(base_dir): return tasks
    niches = sorted(
        d for d in os.listdir(base_dir) if os.path.isdir(os.path.join(base_dir, d)) and not d.startswith("."))
    for niche in niches:
        np = os.path.join(base_dir, niche)
        try:
            strats = sorted(d for d in os.listdir(np) if os.path.isdir(os.path.join(np, d)) and re.match(r"^\d+", d))
        except PermissionError:
            continue
        for s in strats:
            sp = os.path.join(np, s)
            xlsx = next((os.path.join(sp, f) for f in os.listdir(sp) if f.endswith(".xlsx") and not f.startswith("~$")),
                        None)
            tasks.append({"niche": niche, "strat": s, "path": sp, "xlsx": xlsx})
    return tasks


def check_xlsx_filled(xlsx_path):
    if not xlsx_path: return False
    try:
        wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
        ws = wb.active
        filled = sum(1 for row in range(START_ROW, min(START_ROW + 3, END_ROW + 1))
                     if
                     ws.cell(row=row, column=COL_TITLE).value and ws.cell(row=row, column=COL_DESC).value and ws.cell(
                         row=row, column=COL_CTA).value)
        wb.close()
        return filled >= 3
    except:
        return False


def read_unique_xlsx_data(xlsx_path):
    if not xlsx_path or not os.path.exists(xlsx_path): return "", "", ""
    titles, descs, ctas = [], [], []
    seen_t, seen_d, seen_c = set(), set(), set()
    try:
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
        ws = wb.active
        max_search_row = min(START_ROW + 500, ws.max_row + 1)
        for row in range(START_ROW, max_search_row):
            d_val = ws.cell(row=row, column=COL_DESC).value
            c_val = ws.cell(row=row, column=COL_CTA).value
            t_val = ws.cell(row=row, column=COL_TITLE).value
            if t_val and str(t_val).strip() not in seen_t:
                titles.append(str(t_val).strip())
                seen_t.add(str(t_val).strip())
            if d_val and str(d_val).strip() not in seen_d:
                descs.append(str(d_val).strip())
                seen_d.add(str(d_val).strip())
            if c_val and str(c_val).strip() not in seen_c:
                ctas.append(str(c_val).strip())
                seen_c.add(str(c_val).strip())
        wb.close()
        return "\n".join(titles), "\n".join(descs), "\n".join(ctas)
    except:
        return "", "", ""


def update_text_areas(task_idx):
    task = st.session_state.tasks[task_idx]
    t, d, c = read_unique_xlsx_data(task["xlsx"])
    st.session_state["ta_titles"] = t
    st.session_state["ta_descs"] = d
    st.session_state["ta_ctas"] = c


def build_prompt(strat_path, base_dir, target="ALL"):
    for_main = os.path.join(strat_path, "For Main")
    link = read_txt(os.path.join(base_dir, "linksite.txt"))
    tags = read_txt(os.path.join(for_main, "Boards_Tags.txt"))
    board = read_txt(os.path.join(for_main, "Boards_Name.txt"))

    missing = []
    if not link: missing.append("linksite.txt")
    if not tags: missing.append("Boards_Tags.txt")
    if not board: missing.append("Boards_Name.txt")

    json_file = os.path.join(base_dir, "all_products.json")
    if not os.path.exists(json_file):
        json_file = os.path.join(strat_path, "all_products.json")
        if not os.path.exists(json_file): json_file = None

    if missing: return "", missing, 0, None

    if target == "ALL":
        tasks_text = (
            f"1) 100 ЗАГОЛОВКОВ (без эмодзи, без знаков препинания в конце, 3-8 слов, содержат ключевые слова из тегов естественно вписанные)\n"
            f"2) 100 ОПИСАНИЙ (с разными эмодзи по смыслу к содержанию и тегам в начале предложения, без артикля \"A\" в начале, без призыва к действию внутри, 1-2 предложения, должны подходить к ЛЮБОЙ картинке из выдачи по данным тегам)\n"
            f"3) 100 ПРИЗЫВОВ К ДЕЙСТВИЮ — CTA (с разными эмодзи по смыслу, одно короткое предложение, призыв перейти скачать наш товар бесплатно на сайте с подборкой лучшего товара, разные формулировки — не повторяться)\n"
        )
        expected = 3
    elif target == "TITLES":
        tasks_text = f"1) 100 ЗАГОЛОВКОВ (без эмодзи, без знаков препинания в конце, 3-8 слов, содержат ключевые слова из тегов естественно вписанные)\n"
        expected = 1
    elif target == "DESCS":
        tasks_text = f"1) 100 ОПИСАНИЙ (с разными эмодзи по смыслу к содержанию и тегам в начале предложения, без артикля \"A\" в начале, без призыва к действию внутри, 1-2 предложения, должны подходить к ЛЮБОЙ картинке из выдачи по данным тегам)\n"
        expected = 1
    elif target == "CTAS":
        tasks_text = f"1) 100 ПРИЗЫВОВ К ДЕЙСТВИЮ — CTA (с разными эмодзи по смыслу, одно короткое предложение, призыв перейти скачать наш товар бесплатно на сайте с подборкой лучшего товара, разные формулировки — не повторяться)\n"
        expected = 1

    prompt = (
        f"Привет. Я продвигаю на Pinterest товар.\n"
        f"Я прикрепил файл json в котором содержится информация по нашей подборке товара на сайте (детально изучи данный файл, чтобы понять что пользователь получит когда перейдет на данный сайт с фото Pinterest)\n"
        f"Теги по которым продвигаюсь (используй их как контекст для ключевых слов, не вставляй теги дословно в тексты):\n"
        f"{tags}\n"
        f"Главный тег доски: {board}\n\n"
        f"Сделай на английском:\n"
        f"{tasks_text}\n"
        f"ВАЖНЫЕ ПРАВИЛА:\n"
        f"- Любой заголовок + любое описание + любой CTA должны сочетаться друг с другом\n"
        f"- Все тексты в рамках тематики тегов и товара, без выхода за рамки\n"
        f"- Каждый список помести в отдельный блок кода\n"
        f"- Без нумерации, без пустых строк между строками\n"
        f"- В CTA добаявляй с ЗАГЛАВНО БУКВЫ акцент (прмер не fonts, а Fonts) на выгодность перехода на наш сайт, то есть помечай что они получать при переходе на наш сайт по имени товара\n"
        f"- Так же в CTA пиши в 70% случаев FREE (с акцентом на этом и большими буквами), так они могут получить сейчас уже бесплатно то что им нужно - сделай акцент на FREE с большой буквы и чтобы видно было\n"
        f"- (ПРАВИЛО КАСАТЕЛЬНО CTA блок про - ПРОГРАММНОГО ОБЕСПЕЧЕНИЯ/ИНСТРУМЕНТОВ — КРИТИЧЕСКИ ВАЖНО): Если в тегах упоминается ПО (Canva, Adobe, Procreate, Cricut, GoodNotes и т. д.), НЕ пишите фразы определнных фирм программно обеспечения в CTA вроде «Fonts for Canva» или «Typography Adobe». "
        f" Вместо этого делайте акцент на преимуществах, исходя из темы нашей подборки тегов и товаров по ссылке выше. Используйте такие формулировки, которые *идеально подходят* для этих инструментов. Найдите «золотую середину» между тегами и самим продуктом.\n"
        f"- Только сами тексты, ничего лишнего\n"
        f"- Так же сделай здесь чтобы не было подвязки под определенный вид тега (потому что нужно чтобы рандомное фото с рандомного тега взяли и оно подошло под эту нишу)"
    )
    return prompt, missing, expected, json_file


def write_xlsx(xlsx_path, titles, descs, ctas):
    if not titles or not descs or not ctas:
        raise ValueError("Одно из полей пустое! Невозможно записать в Excel.")
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active
    total = END_ROW - START_ROW + 1
    for i in range(total):
        row = START_ROW + i
        ws.cell(row=row, column=COL_DESC).value = descs[i % len(descs)]
        ws.cell(row=row, column=COL_CTA).value = ctas[i % len(ctas)]
        ws.cell(row=row, column=COL_TITLE).value = titles[i % len(titles)]
    wb.save(xlsx_path)


# ─── ИНИЦИАЛИЗАЦИЯ СОСТОЯНИЯ ──────────────────────────────────────────────────
base_dir = os.path.dirname(os.path.abspath(__file__))
if "tasks" not in st.session_state: st.session_state.tasks = scan_tasks(base_dir)
if "done_set" not in st.session_state:
    st.session_state.done_set = {i for i, t in enumerate(st.session_state.tasks) if check_xlsx_filled(t["xlsx"])}
if "cur_idx" not in st.session_state:
    remain = sorted(set(range(len(st.session_state.tasks))) - st.session_state.done_set)
    st.session_state.cur_idx = remain[0] if remain else 0

if "selected_tasks" not in st.session_state:
    st.session_state.selected_tasks = {i for i in range(len(st.session_state.tasks)) if
                                       i not in st.session_state.done_set}

if "qc_state" not in st.session_state:
    st.session_state.qc_state = {i: {"TITLES": "wait", "DESCS": "wait", "CTAS": "wait"} for i in
                                 range(len(st.session_state.tasks))}

if "ta_titles" not in st.session_state: update_text_areas(st.session_state.cur_idx)

tasks = st.session_state.tasks
cur = st.session_state.cur_idx


# ─── ЛОГИКА PLAYWRIGHT ДЛЯ UI ─────────────────────────────────────────────────
def run_ai_for_current_task(task_index, target="ALL"):
    task = tasks[task_index]
    prompt, missing, expected_blocks, json_file = build_prompt(task["path"], base_dir, target)
    if missing:
        st.error(f"Не хватает файлов: {', '.join(missing)}")
        return False

    blocks = interact_with_ai_studio(prompt, expected_blocks, json_file)
    if not blocks or len(blocks) < expected_blocks:
        st.error(f"Сбой парсинга блоков для {task['strat']}.")
        return False

    if target == "ALL":
        st.session_state["ta_titles"] = blocks[0]
        st.session_state["ta_descs"] = blocks[1]
        st.session_state["ta_ctas"] = blocks[2]
    elif target == "TITLES":
        st.session_state["ta_titles"] = blocks[0]
    elif target == "DESCS":
        st.session_state["ta_descs"] = blocks[0]
    elif target == "CTAS":
        st.session_state["ta_ctas"] = blocks[0]
    return True


# ─── МАССОВЫЙ АВТО-ПРОГОН БЕЗ ПАУЗ ───────────────────────────────────────────
if st.session_state.get("auto_run_selected", False):
    remain = sorted([i for i in st.session_state.selected_tasks if i not in st.session_state.done_set])
    if not remain:
        st.success("🎉 Массовый авто-прогон полностью завершен! Все Excel заполнены.")
        st.session_state.auto_run_selected = False
        st.balloons()
    else:
        cur_target = remain[0]
        st.session_state.cur_idx = cur_target
        st.toast(f"🤖 Авто-генерация [{cur_target + 1}/{len(tasks)}]: {tasks[cur_target]['strat']}")

        if run_ai_for_current_task(cur_target, "ALL"):
            t = clean_lines(st.session_state["ta_titles"])
            d = clean_lines(st.session_state["ta_descs"])
            c = clean_lines(st.session_state["ta_ctas"])
            if t and d and c:
                write_xlsx(tasks[cur_target]["xlsx"], t, d, c)
                st.session_state.done_set.add(cur_target)
                st.session_state.qc_state[cur_target] = {"TITLES": "wait", "DESCS": "wait", "CTAS": "wait"}
            st.rerun()  # Мгновенный переход к следующей папке!
        else:
            st.error(f"Сбой генерации для {tasks[cur_target]['strat']}. Авто-прогон остановлен.")
            st.session_state.auto_run_selected = False

# QC-АВТОПРОГОН ДЛЯ ИСПРАВЛЕНИЙ БЕЗ ПАУЗ
elif st.session_state.get("qc_run_redos", False):
    target_idx = None
    targets_to_redo = []
    for i in range(len(tasks)):
        q = st.session_state.qc_state[i]
        if "redo" in q.values():
            target_idx = i
            if q["TITLES"] == "redo": targets_to_redo.append("TITLES")
            if q["DESCS"] == "redo": targets_to_redo.append("DESCS")
            if q["CTAS"] == "redo": targets_to_redo.append("CTAS")
            break

    if target_idx is None:
        st.success("🎉 Все исправления по Контролю Качества завершены!")
        st.session_state.qc_run_redos = False
        st.balloons()
    else:
        st.session_state.cur_idx = target_idx
        update_text_areas(target_idx)
        success = True
        for t in targets_to_redo:
            st.toast(f"🔧 QC исправление {t}: {tasks[target_idx]['strat']}")
            if run_ai_for_current_task(target_idx, t):
                st.session_state.qc_state[target_idx][t] = "ok"  # Помечаем исправленным
            else:
                success = False
                break

        if success:
            t = clean_lines(st.session_state["ta_titles"])
            d = clean_lines(st.session_state["ta_descs"])
            c = clean_lines(st.session_state["ta_ctas"])
            write_xlsx(tasks[target_idx]["xlsx"], t, d, c)
            st.session_state.done_set.add(target_idx)
            st.rerun()  # Мгновенно переходим к следующему исправлению!
        else:
            st.session_state.qc_run_redos = False

# ─── SIDEBAR ──────────────────────────────────────────────────────────────────
with st.sidebar:
    st.markdown('<div class="main-header">🤖 Auto Filler</div>', unsafe_allow_html=True)
    if tasks: st.progress(len(st.session_state.done_set) / len(tasks),
                          text=f"Заполнено Excel: {len(st.session_state.done_set)} / {len(tasks)}")

    st.markdown("### 📋 Список папок")
    c_sel1, c_sel2 = st.columns(2)
    if c_sel1.button("✅ Выбрать все", use_container_width=True):
        st.session_state.selected_tasks = set(range(len(tasks)))
        st.rerun()
    if c_sel2.button("❌ Снять все", use_container_width=True):
        st.session_state.selected_tasks = set()
        st.rerun()

    st.caption("Галочка = папка участвует в Массовых действиях")

    for i, t in enumerate(tasks):
        col_cb, col_btn = st.columns([1, 6])
        is_sel = i in st.session_state.selected_tasks

        if col_cb.checkbox("Выбрать", value=is_sel, key=f"sel_{i}", label_visibility="collapsed") != is_sel:
            if is_sel:
                st.session_state.selected_tasks.remove(i)
            else:
                st.session_state.selected_tasks.add(i)
            st.rerun()

        icon = "✅" if i in st.session_state.done_set else ("📊" if t["xlsx"] else "❌")
        if i == cur:
            col_btn.markdown(
                f"<div style='background:#1e1030;border:1px solid #a78bfa;border-radius:6px;padding:4px;font-size:11px;color:#c4b5fd;'>▶ {t['strat']}</div>",
                unsafe_allow_html=True)
        else:
            if col_btn.button(f"{icon} {t['strat']}", key=f"nav_{i}", use_container_width=True):
                st.session_state.cur_idx = i
                update_text_areas(i)
                st.rerun()

# ─── MAIN APP ─────────────────────────────────────────────────────────────────
if not tasks: st.stop()
task = tasks[cur]

tab_gen, tab_qc = st.tabs(["🚀 Режим Генерации & Управления", "👁 Режим QC (Контроль Качества)"])

# ==============================================================================
# ВКЛАДКА 1: ГЕНЕРАЦИЯ
# ==============================================================================
with tab_gen:
    st.markdown("### ⚡ Массовые действия (без пауз)")
    c1, c2 = st.columns(2)
    with c1:
        if st.button("▶️ Заполнить ПУСТЫЕ выбранные (Автоматом)", type="primary", use_container_width=True):
            st.session_state.auto_run_selected = True
            st.rerun()
    with c2:
        if st.button("⚠️ ПЕРЕЗАПИСАТЬ выбранные (Заново)", use_container_width=True):
            st.session_state.show_confirm_overwrite = True
            st.rerun()

    if st.session_state.get("show_confirm_overwrite", False):
        st.warning(
            f"Начнется НЕПРЕРЫВНЫЙ авто-прогон для {len(st.session_state.selected_tasks)} папок. Старые данные в Excel будут перезаписаны сразу по мере генерации.")
        cw1, cw2 = st.columns(2)
        if cw1.button("🔴 ДА, НАЧАТЬ МАССОВУЮ ПЕРЕЗАПИСЬ", type="primary"):
            st.session_state.show_confirm_overwrite = False
            st.session_state.done_set = st.session_state.done_set - st.session_state.selected_tasks
            st.session_state.auto_run_selected = True
            st.rerun()
        if cw2.button("Отмена"):
            st.session_state.show_confirm_overwrite = False
            st.rerun()

    st.divider()

    json_file_path = os.path.join(base_dir, "all_products.json")
    if not os.path.exists(json_file_path): json_file_path = os.path.join(task['path'], "all_products.json")
    json_text = f"📎 JSON: {os.path.basename(json_file_path)}" if os.path.exists(
        json_file_path) else "⚠️ JSON не найден!"

    st.markdown(f"""
    <div class="card card-accent">
      <div class="niche-badge">📁 {task['niche']} ➜ {task['strat']}</div>
      <div style="font-size: 11px; color: #34d399;">📊 {task['xlsx']}</div>
      <div style="font-size: 11px; color: #a78bfa; margin-top: 3px;">{json_text}</div>
    </div>
    """, unsafe_allow_html=True)

    mc1, mc2 = st.columns(2)
    with mc1:
        if st.button("🤖 Сгенерировать ВСЁ для этой папки", use_container_width=True):
            with st.spinner("Работаем с AI..."):
                if run_ai_for_current_task(cur, "ALL"): st.rerun()
    with mc2:
        with st.popover("🔄 Точечная регенерация (исправить)"):
            if st.button("Переделать ЗАГОЛОВКИ"):
                if run_ai_for_current_task(cur, "TITLES"): st.rerun()
            if st.button("Переделать ОПИСАНИЯ"):
                if run_ai_for_current_task(cur, "DESCS"): st.rerun()
            if st.button("Переделать CTA"):
                if run_ai_for_current_task(cur, "CTAS"): st.rerun()

    col1, col2, col3 = st.columns(3)
    with col1:
        raw_t = st.text_area("📝 ЗАГОЛОВКИ", height=340, key="ta_titles")
        titles = clean_lines(raw_t)
        st.caption(f"Строк: {len(titles)}")
    with col2:
        raw_d = st.text_area("📖 ОПИСАНИЯ", height=340, key="ta_descs")
        descs = clean_lines(raw_d)
        st.caption(f"Строк: {len(descs)}")
    with col3:
        raw_c = st.text_area("🎯 CTA", height=340, key="ta_ctas")
        ctas = clean_lines(raw_c)
        st.caption(f"Строк: {len(ctas)}")

    # ─── ФИНАЛЬНАЯ РУЧНАЯ ЗАПИСЬ ──────────────────────────────────────────────
    write_ready = bool(titles and descs and ctas and task["xlsx"])
    if st.button("💾 СОХРАНИТЬ ТЕКУЩУЮ В EXCEL", type="primary", use_container_width=True, disabled=not write_ready):
        write_xlsx(task["xlsx"], titles, descs, ctas)
        st.session_state.done_set.add(cur)
        st.session_state.qc_state[cur] = {"TITLES": "wait", "DESCS": "wait", "CTAS": "wait"}
        st.success("✅ Сохранено!")

# ==============================================================================
# ВКЛАДКА 2: QC (Контроль Качества)
# ==============================================================================
with tab_qc:
    st.markdown("### 👁 Контроль качества заполнения")
    st.caption("Проверьте тексты из Excel. Пометьте, какие блоки подходят, а какие ИИ должен переписать.")

    qc_t = st.session_state.qc_state[cur]["TITLES"]
    qc_d = st.session_state.qc_state[cur]["DESCS"]
    qc_c = st.session_state.qc_state[cur]["CTAS"]


    def set_qc(t, d, c):
        st.session_state.qc_state[cur] = {"TITLES": t, "DESCS": d, "CTAS": c}


    qc_b1, qc_b2 = st.columns(2)
    if qc_b1.button("✅ ВСЁ ОТЛИЧНО (В этой папке)", use_container_width=True):
        set_qc("ok", "ok", "ok")
        st.rerun()
    if qc_b2.button("🔴 ПЕРЕДЕЛАТЬ ВСЁ (В этой папке)", use_container_width=True):
        set_qc("redo", "redo", "redo")
        st.rerun()

    st.divider()


    def qc_radio(label, current_val, key):
        opts = {"wait": "⏳ Ждет проверки", "ok": "🟢 ОК", "redo": "🔴 Переделать"}
        return st.radio(label, options=list(opts.keys()), format_func=lambda x: opts[x],
                        index=list(opts.keys()).index(current_val), key=key, horizontal=True)


    col_q1, col_q2, col_q3 = st.columns(3)
    with col_q1:
        new_t = qc_radio("Оценка ЗАГОЛОВКОВ", qc_t, f"qct_{cur}")
        st.info("\n\n".join(clean_lines(st.session_state.get("ta_titles", ""))[:5]) + "\n\n...(показаны первые 5)")
    with col_q2:
        new_d = qc_radio("Оценка ОПИСАНИЙ", qc_d, f"qcd_{cur}")
        st.success("\n\n".join(clean_lines(st.session_state.get("ta_descs", ""))[:5]) + "\n\n...(показаны первые 5)")
    with col_q3:
        new_c = qc_radio("Оценка CTA", qc_c, f"qcc_{cur}")
        st.warning("\n\n".join(clean_lines(st.session_state.get("ta_ctas", ""))[:5]) + "\n\n...(показаны первые 5)")

    if new_t != qc_t or new_d != qc_d or new_c != qc_c:
        set_qc(new_t, new_d, new_c)
        st.rerun()

    st.divider()

    total_redos = sum(1 for q in st.session_state.qc_state.values() for v in q.values() if v == "redo")
    st.markdown(f"**Блоков, отправленных на переделку (по всем папкам):** `{total_redos}`")

    if total_redos > 0:
        if st.button("🚀 ЗАВЕРШИТЬ КОНТРОЛЬ И ИСПРАВИТЬ БРАК (Автоматом без пауз)", type="primary",
                     use_container_width=True):
            st.session_state.show_confirm_qc = True
            st.rerun()

        if st.session_state.get("show_confirm_qc", False):
            st.warning(f"Начнется непрерывный авто-прогон для {total_redos} бракованных блоков. Запустить?")
            c_qc1, c_qc2 = st.columns(2)
            if c_qc1.button("✅ ДА, ИСПРАВИТЬ ВСЁ"):
                st.session_state.show_confirm_qc = False
                st.session_state.qc_run_redos = True
                st.rerun()
            if c_qc2.button("❌ Отмена"):
                st.session_state.show_confirm_qc = False
                st.rerun()