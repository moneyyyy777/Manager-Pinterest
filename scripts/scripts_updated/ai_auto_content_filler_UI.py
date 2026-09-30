import sys
import subprocess
import os

# 1. ПРОВЕРЯЕМ АВТОЗАПУСК СТРОГО ДО ИМПОРТА STREAMLIT!
if "streamlit" not in sys.modules:
    script = os.path.abspath(__file__)
    subprocess.run([sys.executable, "-m", "streamlit", "run", script])
    sys.exit()

# 2. ИМПОРТИРУЕМ ВСЁ ОСТАЛЬНОЕ
import json
import time
import re
import openpyxl
import streamlit as st
from playwright.sync_api import sync_playwright, TimeoutError

# ─── НАСТРОЙКИ ────────────────────────────────────────────────────────────────
START_ROW = 2
END_ROW = 5001
COL_DESC = 12
COL_CTA = 13
COL_TITLE = 14
REQUIRED_LINES = 100

CHROME_DEBUG_URL = "http://127.0.0.1:9222"

st.set_page_config(page_title="🤖 Pinterest Filler & QC", layout="wide", initial_sidebar_state="expanded")

st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=JetBrains+Mono:wght@400;700&family=Syne:wght@700;800&display=swap');
html, body, [class*="css"] { background-color: #0d0d14 !important; color: #e0e0f0 !important; font-family: 'JetBrains Mono', monospace !important; }
section[data-testid="stSidebar"] { background: #111119 !important; border-right: 1px solid #1e1e2e !important; }

/* ─── ТОЧЕЧНЫЙ БЕЛЫЙ ЦВЕТ ДЛЯ 3 ЭЛЕМЕНТОВ В САЙДБАРЕ ─── */
section[data-testid="stSidebar"] h3 { color: #ffffff !important; }
section[data-testid="stSidebar"] .stProgress p { color: #ffffff !important; }
section[data-testid="stSidebar"] [data-testid="stCaptionContainer"] p { color: #ffffff !important; }

.niche-badge {
    display: inline-block; 
    background: #25183e !important;
    border: 1px solid #a78bfa !important; 
    border-radius: 6px;
    padding: 6px 14px; 
    font-size: 13px; 
    font-weight: 700;
    color: #ffffff !important; 
    margin-bottom: 8px;
}
.main-header { font-family: 'Syne', sans-serif; font-size: 26px; font-weight: 800; background: linear-gradient(90deg, #a78bfa, #f472b6); -webkit-background-clip: text; -webkit-text-fill-color: transparent; }
.card { background: #13131e; border: 1px solid #1e1e30; border-radius: 10px; padding: 16px 20px; margin-bottom: 12px; }
.card-accent { border-left: 3px solid #a78bfa; }
div[data-testid="stTextArea"] textarea { background: #0a0a12 !important; color: #c0c0e0 !important; border: 1px solid #252538 !important; }
div.stButton > button[kind="primary"] { background: linear-gradient(135deg, #7c3aed, #a855f7) !important; border: none !important; color: white !important; font-weight: 700 !important; }
.stRadio > div { background: #1a1a2e; padding: 5px 10px; border-radius: 8px; border: 1px solid #2e2e4e; margin-bottom: 10px; }
.stRadio * { color: #ffffff !important; font-size: 12px !important; }
</style>
""", unsafe_allow_html=True)


# ─── ИНТЕГРАЦИЯ С PLAYWRIGHT (AI STUDIO) ──────────────────────────────────────
def interact_with_ai_studio(prompt: str, expected_blocks: int, json_file_path: str, max_attempts: int = 15):
    with sync_playwright() as p:
        try:
            browser = p.chromium.connect_over_cdp(CHROME_DEBUG_URL)
            context = browser.contexts[0]

            page = None
            for p_ext in context.pages:
                if "aistudio" in p_ext.url:
                    page = p_ext
                    break
            if not page:
                page = context.new_page()

            for attempt in range(1, max_attempts + 1):
                try:
                    page.bring_to_front()
                    st.toast(f"🔄 Попытка {attempt}/{max_attempts}: Подготовка AI Studio...")

                    page.goto("https://aistudio.google.com/app/u/1/prompts/new_chat", timeout=30000)
                    page.wait_for_timeout(2000)

                    if json_file_path and os.path.exists(json_file_path):
                        try:
                            file_input = page.locator('input[type="file"]').first
                            file_input.set_input_files(json_file_path)
                            page.wait_for_timeout(3000)
                        except Exception as e:
                            print(f"⚠️ Ошибка загрузки файла: {e}")

                    text_area = page.locator("textarea").last
                    text_area.fill(prompt)
                    page.wait_for_timeout(500)

                    try:
                        grounding_sw = page.locator('button[role="switch"][aria-label*="Grounding with Google Search"]')
                        if grounding_sw.is_visible(timeout=500) and grounding_sw.get_attribute(
                                "aria-checked") == "true":
                            grounding_sw.click()
                            page.wait_for_timeout(300)

                        url_sw = page.locator('button[role="switch"][aria-label*="Browse the url context"]')
                        if url_sw.is_visible(timeout=500) and url_sw.get_attribute("aria-checked") == "true":
                            url_sw.click()
                            page.wait_for_timeout(300)
                    except Exception:
                        pass

                    try:
                        clicked = page.evaluate('''() => {
                            const btn = document.querySelector('button[mattooltipclass="run-button-tooltip"]') 
                                     || document.querySelector('button.ctrl-enter-submits');
                            if (btn) { btn.click(); return true; }
                            return false;
                        }''')
                        if not clicked:
                            page.locator('button[mattooltipclass="run-button-tooltip"]').first.click(force=True)
                    except Exception:
                        text_area.focus()
                        page.keyboard.press("Meta+Enter")

                    st.toast(f"⏳ Попытка {attempt}: AI генерирует ответ...")

                    prev_len = 0
                    stable_ticks = 0

                    for _ in range(90):  # До 3 минут генерации
                        time.sleep(2)
                        page_content = page.content()
                        if "An internal error has occurred" in page_content or "Failed to generate content" in page_content:
                            raise Exception("AI Studio выдал внутреннюю ошибку.")

                        code_blocks = page.locator("ms-code-block pre").all_inner_texts()
                        if not code_blocks:
                            code_blocks = page.locator("pre").all_inner_texts()

                        if len(code_blocks) >= expected_blocks:
                            curr_len = len(code_blocks[-1])
                            if curr_len > 30 and curr_len == prev_len:
                                stable_ticks += 1
                                if stable_ticks >= 3:
                                    final_blocks = page.locator("ms-code-block pre").all_inner_texts()
                                    if not final_blocks:
                                        final_blocks = page.locator("pre").all_inner_texts()
                                    if len(final_blocks) != expected_blocks:
                                        raise Exception(
                                            f"AI вернул {len(final_blocks)} блоков кода вместо ожидаемых {expected_blocks} "
                                            f"— результат ненадёжен, переспрашиваю заново."
                                        )
                                    return final_blocks[-expected_blocks:]
                            else:
                                stable_ticks = 0
                                prev_len = curr_len

                    raise Exception("Таймаут ожидания ответа AI.")

                except Exception as attempt_err:
                    st.warning(f"⚠️ Попытка {attempt} сбой: {attempt_err}. Перезапуск через 5 сек...")
                    time.sleep(5)

            raise Exception(f"Все {max_attempts} попыток завершились ошибкой AI Studio.")
        except Exception as e:
            st.error(f"❌ Ошибка Playwright: {e}")
            return []


# ─── ХЕЛПЕРЫ ──────────────────────────────────────────────────────────────────
def read_txt(path):
    if not os.path.exists(path): return ""
    with open(path, "r", encoding="utf-8", errors="replace") as f: return f.read().strip()


def clean_lines(raw):
    if not raw: return []
    lines = []
    for line in raw.splitlines():
        line = line.strip()
        line = re.sub(r'^\d+[\.\)]\s*|^[*\-•]\s*|^`{3,}', '', line).replace("```", "").strip()
        if line: lines.append(line)
    return lines


def load_qc_state(task_path):
    """Читает состояние QC из файла qc_state.json в папке задачи."""
    json_path = os.path.join(task_path, "qc_state.json")
    if os.path.exists(json_path):
        try:
            with open(json_path, "r", encoding="utf-8") as f:
                return json.load(f)
        except:
            pass
    return {"TITLES": "wait", "DESCS": "wait", "CTAS": "wait"}


def save_qc_state(task_path, state):
    """Сохраняет состояние QC в файл qc_state.json в папку задачи."""
    json_path = os.path.join(task_path, "qc_state.json")
    try:
        with open(json_path, "w", encoding="utf-8") as f:
            json.dump(state, f, ensure_ascii=False, indent=2)
    except Exception as e:
        print(f"Ошибка сохранения QC: {e}")


def scan_tasks(base_dir):
    tasks = []
    if not os.path.isdir(base_dir): return tasks
    niches = sorted(
        d for d in os.listdir(base_dir) if os.path.isdir(os.path.join(base_dir, d)) and not d.startswith("."))
    for niche in niches:
        np = os.path.join(base_dir, niche)
        try:
            strats = sorted(d for d in os.listdir(np) if os.path.isdir(os.path.join(np, d)) and re.match(r"^\d+", d))
        except PermissionError:
            continue
        for s in strats:
            sp = os.path.join(np, s)
            xlsx = next((os.path.join(sp, f) for f in os.listdir(sp) if f.endswith(".xlsx") and not f.startswith("~$")),
                        None)
            tasks.append({"niche": niche, "strat": s, "path": sp, "xlsx": xlsx})
    return tasks


def check_xlsx_filled(xlsx_path):
    if not xlsx_path: return False
    try:
        wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
        ws = wb.active
        filled = sum(1 for row in range(START_ROW, min(START_ROW + 3, END_ROW + 1))
                     if
                     ws.cell(row=row, column=COL_TITLE).value and ws.cell(row=row, column=COL_DESC).value and ws.cell(
                         row=row, column=COL_CTA).value)
        wb.close()
        return filled >= 3
    except:
        return False


def read_unique_xlsx_data(xlsx_path):
    if not xlsx_path or not os.path.exists(xlsx_path): return "", "", ""
    titles, descs, ctas = [], [], []
    seen_t, seen_d, seen_c = set(), set(), set()
    try:
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
        ws = wb.active
        max_search_row = min(START_ROW + 500, ws.max_row + 1)
        for row in range(START_ROW, max_search_row):
            d_val = ws.cell(row=row, column=COL_DESC).value
            c_val = ws.cell(row=row, column=COL_CTA).value
            t_val = ws.cell(row=row, column=COL_TITLE).value
            if t_val and str(t_val).strip() not in seen_t:
                titles.append(str(t_val).strip())
                seen_t.add(str(t_val).strip())
            if d_val and str(d_val).strip() not in seen_d:
                descs.append(str(d_val).strip())
                seen_d.add(str(d_val).strip())
            if c_val and str(c_val).strip() not in seen_c:
                ctas.append(str(c_val).strip())
                seen_c.add(str(c_val).strip())
        wb.close()
        return "\n".join(titles), "\n".join(descs), "\n".join(ctas)
    except:
        return "", "", ""


def update_text_areas(task_idx):
    task = st.session_state.tasks[task_idx]
    t, d, c = read_unique_xlsx_data(task["xlsx"])
    st.session_state["ta_titles"] = t
    st.session_state["ta_descs"] = d
    st.session_state["ta_ctas"] = c


def build_prompt(strat_path, base_dir, target="ALL", niche_mode=False):
    for_main = os.path.join(strat_path, "For Main")
    link = read_txt(os.path.join(base_dir, "linksite.txt"))
    tags = read_txt(os.path.join(for_main, "Boards_Tags.txt"))
    board = read_txt(os.path.join(for_main, "Boards_Name.txt"))

    missing = []
    if not link: missing.append("linksite.txt")
    if not tags: missing.append("Boards_Tags.txt")
    if not board: missing.append("Boards_Name.txt")

    json_file = os.path.join(base_dir, "all_products.json")
    if not os.path.exists(json_file):
        json_file = os.path.join(strat_path, "all_products.json")
        if not os.path.exists(json_file): json_file = None

    if missing: return "", missing, 0, None, []

    # ── Описание каждого блока задачи. Позволяет собрать промпт под ЛЮБОЙ набор
    # блоков (1, 2 или 3 сразу) одним проходом — вместо трёх отдельных промптов. ──
    TASK_LABELS = {
        "TITLES": (
            "РОВНО 100 ЗАГОЛОВКОВ (не 99, не 80, не 105 — точно 100 строк)",
            "без эмодзи, без знаков препинания в конце, 3-8 слов, содержат ключевые слова из тегов естественно вписанные"
        ),
        "DESCS": (
            "РОВНО 100 ОПИСАНИЙ (не 99, не 80, не 105 — точно 100 строк)",
            "⚠️ ОБЯЗАТЕЛЬНО в КАЖДОЙ строке ставь эмодзи по смыслу к содержанию и тегам в самом начале предложения — "
            "это требование не снимается ни при каких условиях, без артикля \"A\" в начале, без призыва к действию внутри, "
            "1-2 предложения, должны подходить к ЛЮБОЙ картинке из выдачи по данным тегам"
        ),
        "CTAS": (
            "РОВНО 100 ПРИЗЫВОВ К ДЕЙСТВИЮ — CTA (не 99, не 80, не 105 — точно 100 строк)",
            "⚠️ ОБЯЗАТЕЛЬНО в КАЖДОЙ строке ставь эмодзи по смыслу в самом начале — "
            "это требование не снимается ни при каких условиях, одно короткое предложение, призыв перейти скачать наш товар "
            "бесплатно на сайте с подборкой лучшего товара, разные формулировки — не повторяться"
        ),
    }
    CANONICAL_ORDER = ["TITLES", "DESCS", "CTAS"]

    # target может быть "ALL", одиночной строкой ("TITLES"/"DESCS"/"CTAS")
    # или списком/множеством нескольких ключей сразу — например ["TITLES", "CTAS"].
    if target == "ALL":
        blocks_order = CANONICAL_ORDER[:]
    elif isinstance(target, str):
        blocks_order = [target]
    else:
        blocks_order = [k for k in CANONICAL_ORDER if k in target]

    tasks_text = ""
    for idx, key in enumerate(blocks_order, start=1):
        name, desc = TASK_LABELS[key]
        tasks_text += f"{idx}) {name} ({desc})\n"
    expected = len(blocks_order)

    # ── Блок "ОБЩАЯ НИША" — включается тумблером в UI, когда обычная генерация
    # промахивается мимо части фото (например теги смешивают еду и напитки,
    # а тексты получаются заточены только под "десерты"). ──
    niche_block = ""
    if niche_mode:
        niche_block = (
            f"ШАГ 0 — ОПРЕДЕЛИ ОБЩУЮ НИШУ (сделай это ПЕРЕД генерацией текстов):\n"
            f"Внимательно прочитай ВЕСЬ список тегов выше целиком, от первого до последнего — не цепляйся за первый попавшийся тег "
            f"и не суди по 2-3 тегам.\n\n"
            f"В нише ВСЕГДА есть два независимых слоя, и оба нужно достать ИМЕННО ИЗ ТЕГОВ, ничего не придумывая от себя:\n"
            f"  1. ТИП ПРОДУКТА / ТЕХНИКА / ФОРМАТ — то, КАК это сделано или в каком формате продаётся "
            f"(watercolor, gouache, pastel, embroidery, cross stitch, svg cut file, png clipart, digital paper, "
            f"line art, vector, sublimation, printable, font/typeface и т.п.).\n"
            f"  2. ТЕМА / СЮЖЕТ — то, ЧТО изображено или про что это "
            f"(food, drinks, fruits, fashion accessories, wedding, floral, animals и т.п.).\n\n"
            f"ВАЖНО: слой 'ТИП ПРОДУКТА' бери ТОЛЬКО если соответствующее слово реально встречается в списке тегов выше. "
            f"НИКОГДА не додумывай и не подставляй тип продукта по умолчанию, даже если он кажется самым вероятным для этой темы. "
            f"Одна и та же тема 'Food & Drinks' может встречаться в подборках с тегами watercolor clipart, а может — "
            f"в подборках с тегами embroidery designs, cross stitch patterns, line art coloring pages и т.д. — "
            f"это РАЗНЫЕ ниши, и итоговое название ниши должно отражать именно то, что написано в тегах, а не твою догадку.\n\n"
            f"Если в тегах встречается несколько разных слов уровня 'тип продукта' одновременно (например watercolor + gouache + pastel) — "
            f"это признаки одной художественной техники, объедини их в один общий термин уровня ниши (например 'Watercolor').\n"
            f"Если слова уровня 'тип продукта' в тегах вообще нет — не добавляй его в название ниши, оставь только тему/сюжет.\n\n"
            f"Общая ниша = [слой 'тема/сюжет', объединяющий ВСЕ под-темы тегов] + [слой 'тип продукта', ТОЛЬКО если он реально есть в тегах]. Примеры (это только иллюстрация принципа, а не готовые ответы):\n"
            f"  - Теги: watercolor clipart, gouache clipart, pastel clipart, sweets, drinks, fruits "
            f"→ ниша: 'Food & Drinks Watercolor Clipart' (тема Food&Drinks объединяет sweets/drinks/fruits, тип продукта Watercolor взят из тегов).\n"
            f"  - Теги: food embroidery, drinks embroidery designs, fruit embroidery pattern "
            f"→ ниша: 'Food & Drinks Embroidery' (та же тема Food&Drinks, но тип продукта уже Embroidery, а не Watercolor — потому что именно это написано в тегах).\n"
            f"  - Теги: bags, jewelry, shoes, accessories clipart "
            f"→ ниша: 'Fashion Accessories Clipart'.\n"
            f"  - Теги: wedding font, vintage script font, hand lettered font "
            f"→ ниша: 'Elegant Wedding Typography' (тип продукта здесь — сам шрифт, это следует из тегов).\n\n"
            f"Твоя задача — определить САМУЮ ШИРОКУЮ тему, которая ОДИНАКОВО ХОРОШО описывает АБСОЛЮТНО ВСЕ теги из списка, "
            f"а не только часть из них, и честно (без додумывания) добавить к ней тип продукта, если он есть в тегах.\n"
            f"Напиши ПЕРВОЙ строкой ответа (вне блоков кода, до них): НИША: <общая ниша на английском, 2-6 слов>\n\n"
            f"ШАГ 1 — ГЕНЕРАЦИЯ ПОД ЭТУ НИШУ:\n"
            f"Дальше генерируй тексты, ориентируясь ИМЕННО на широкую нишу из Шага 0, а не на отдельные теги. "
            f"Каждый заголовок/описание/CTA должен одинаково хорошо подходить к ЛЮБОЙ картинке, которая может выпасть по ЛЮБОМУ тегу из списка "
            f"(в примере с едой — и к фото десерта, и к фото кофе, и к фото фруктов одновременно), а не только к картинкам одной под-категории. "
            f"Не используй слова, которые жёстко привязывают текст к одной под-теме, если ниша шире (например не пиши 'dessert' или 'sweets', "
            f"если ниша — 'Food & Drinks'; используй нейтральные общие слова уровня ниши: 'treats', 'illustrations', 'graphics', 'designs' и т.п.). "
            f"Точно так же не пиши слово типа продукта, которого не было в определённой нише (если ниша получилась 'Food & Drinks Embroidery' — "
            f"не упоминай 'watercolor', и наоборот).\n\n"
            f"⚠️ ШАГ 0 — ЭТО ДОПОЛНИТЕЛЬНАЯ ЗАДАЧА, А НЕ ЗАМЕНА ОСНОВНОЙ. Она НЕ отменяет и НЕ ослабляет ни одно из требований по формату "
            f"(эмодзи, длина, знаки препинания и т.д.), которые указаны ниже в разделе 'Сделай на английском' и в 'ВАЖНЫЕ ПРАВИЛА'. "
            f"Определение ниши — это только первый шаг перед генерацией, само содержание текстов после этого шага "
            f"должно ТОЧНО так же соблюдать все требования к формату, как будто Шага 0 не было.\n\n"
        )

    prompt = (
        f"Привет. Я продвигаю на Pinterest товар.\n"
        f"Я прикрепил файл json в котором содержится информация по нашей подборке товара на сайте (детально изучи данный файл, чтобы понять что пользователь получит когда перейдет на данный сайт с фото Pinterest)\n"
        f"Теги по которым продвигаюсь (используй их как контекст для ключевых слов, не вставляй теги дословно в тексты):\n"
        f"{tags}\n"
        f"Главный тег доски: {board}\n\n"
f"{niche_block}"
        f"Сделай на английском:\n"
        f"{tasks_text}\n"
        f"⚠️ КРИТИЧЕСКИ ВАЖНО: в ответе должно быть РОВНО {expected} блок(ов) кода — ни больше, ни меньше. "
        f"Не добавляй никаких дополнительных блоков, черновиков, вариантов или примеров сверх перечисленных выше {expected} пункт(ов). "
        f"Если тебя просят только про одну часть (например только заголовки) — НЕ генерируй заодно описания или CTA, даже для полноты.\n"
        f"ВАЖНЫЕ ПРАВИЛА:\n"
        f"- Любой заголовок + любое описание + любой CTA должны сочетаться друг с другом\n"
        f"- Все тексты в рамках тематики тегов и товара, без выхода за рамки\n"
        f"- Каждый список помести в отдельный блок кода\n"
        f"- Без нумерации, без пустых строк между строками\n"
        f"- В CTA добаявляй с ЗАГЛАВНО БУКВЫ акцент (прмер не fonts, а Fonts) на выгодность перехода на наш сайт, то есть помечай что они получать при переходе на наш сайт по имени товара\n"
        f"- Так же в CTA пиши в 70% случаев FREE (с акцентом на этом и большими буквами), так они могут получить сейчас уже бесплатно то что им нужно - сделай акцент на FREE с большой буквы и чтобы видно было\n"
        f"- (ПРАВИЛО КАСАТЕЛЬНО CTA блок про - ПРОГРАММНОГО ОБЕСПЕЧЕНИЯ/ИНСТРУМЕНТОВ — КРИТИЧЕСКИ ВАЖНО): Если в тегах упоминается ПО (Canva, Adobe, Procreate, Cricut, GoodNotes и т. д.), НЕ пишите фразы определнных фирм программно обеспечения в CTA вроде «Fonts for Canva» или «Typography Adobe». "
        f" Вместо этого делайте акцент на преимуществах, исходя из темы нашей подборки тегов и товаров по ссылке выше. Используйте такие формулировки, которые *идеально подходят* для этих инструментов. Найдите «золотую середину» между тегами и самим продуктом.\n"
        f"- Только сами тексты, ничего лишнего\n"
        f"- Так же сделай здесь чтобы не было подвязки под определенный вид тега (потому что нужно чтобы рандомное фото с рандомного тега взяли и оно подошло под эту нишу)\n"
        f"- ⚠️ НАПОМИНАНИЕ В КОНЦЕ (самое важное, не забудь): в ОПИСАНИЯХ и в CTA эмодзи в начале КАЖДОЙ строки — это жёсткое обязательное требование, "
        f"проверь перед выводом, что оно выполнено в каждой без исключения строке этих двух блоков.\n"
        f"- ⚠️ ПЕРЕД ТЕМ КАК ОТПРАВИТЬ ОТВЕТ — посчитай количество строк в каждом блоке кода. В каждом блоке должно быть РОВНО 100 строк. "
        f"Если получилось меньше — дописывай новые строки, пока не наберётся ровно 100. Если больше — удали лишние. Это критическое требование."
)
    return prompt, missing, expected, json_file, blocks_order

def write_xlsx(xlsx_path, titles, descs, ctas):
    if not titles or not descs or not ctas:
        raise ValueError("Одно из полей пустое! Невозможно записать в Excel.")
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active
    total = END_ROW - START_ROW + 1
    for i in range(total):
        row = START_ROW + i
        ws.cell(row=row, column=COL_DESC).value = descs[i % len(descs)]
        ws.cell(row=row, column=COL_CTA).value = ctas[i % len(ctas)]
        ws.cell(row=row, column=COL_TITLE).value = titles[i % len(titles)]
    wb.save(xlsx_path)


# ─── ИНИЦИАЛИЗАЦИЯ СОСТОЯНИЯ ──────────────────────────────────────────────────
base_dir = os.path.dirname(os.path.abspath(__file__))

if "tasks" not in st.session_state:
    st.session_state.tasks = scan_tasks(base_dir)

if "qc_state" not in st.session_state:
    # Загружаем QC стейт из json-файлов для каждой папки
    st.session_state.qc_state = {}
    for i, t in enumerate(st.session_state.tasks):
        st.session_state.qc_state[i] = load_qc_state(t["path"])

        # Предохранитель при (пере)запуске скрипта: если в блоке реально сохранено
        # не ровно REQUIRED_LINES строк — статус ОБЯЗАН быть "redo", даже если
        # в qc_state.json было записано "ok" или "wait".
        if t["xlsx"]:
            tt, dd, cc = read_unique_xlsx_data(t["xlsx"])
            counts = {
                "TITLES": len(clean_lines(tt)),
                "DESCS": len(clean_lines(dd)),
                "CTAS": len(clean_lines(cc)),
            }
            changed = False
            for key, cnt in counts.items():
                if cnt > 0 and cnt != REQUIRED_LINES and st.session_state.qc_state[i][key] != "redo":
                    st.session_state.qc_state[i][key] = "redo"
                    changed = True
            if changed:
                save_qc_state(t["path"], st.session_state.qc_state[i])

if "done_set" not in st.session_state:
    st.session_state.done_set = {i for i, t in enumerate(st.session_state.tasks) if check_xlsx_filled(t["xlsx"])}

if "cur_idx" not in st.session_state:
    remain = sorted(set(range(len(st.session_state.tasks))) - st.session_state.done_set)
    st.session_state.cur_idx = remain[0] if remain else 0

if "selected_tasks" not in st.session_state:
    st.session_state.selected_tasks = {i for i in range(len(st.session_state.tasks)) if
                                       i not in st.session_state.done_set}

if "ta_titles" not in st.session_state:
    update_text_areas(st.session_state.cur_idx)

tasks = st.session_state.tasks
cur = st.session_state.cur_idx


if "niche_mode" not in st.session_state:
    st.session_state.niche_mode = False

if "just_regenerated" not in st.session_state:
    st.session_state.just_regenerated = {}

MAX_LINE_COUNT_RETRIES = 3  # сколько раз пере-промптить блок, если строк меньше 100

# ─── ЛОГИКА PLAYWRIGHT ДЛЯ UI ─────────────────────────────────────────────────
def run_ai_for_current_task(task_index, target="ALL", niche_mode=None):
    task = tasks[task_index]
    if niche_mode is None:
        niche_mode = st.session_state.get("niche_mode", False)
    prompt, missing, expected_blocks, json_file, blocks_order = build_prompt(
        task["path"], base_dir, target, niche_mode=niche_mode
    )
    if missing:
        st.error(f"Не хватает файлов: {', '.join(missing)}")
        return False

    blocks = interact_with_ai_studio(prompt, expected_blocks, json_file)
    if not blocks or len(blocks) < expected_blocks:
        st.error(f"Сбой парсинга блоков для {task['strat']}.")
        return False

    key_to_state = {"TITLES": "ta_titles", "DESCS": "ta_descs", "CTAS": "ta_ctas"}

    # Раскладываем блоки, но сразу проверяем количество строк в каждом.
    for i, key in enumerate(blocks_order):
        lines = clean_lines(blocks[i])

        # Больше 100 — просто обрезаем до 100.
        if len(lines) > REQUIRED_LINES:
            lines = lines[:REQUIRED_LINES]

        # Меньше 100 — пере-промптиваем ИМЕННО этот блок отдельно, пока не наберём 100
        # (или не кончатся попытки).
        retries = 0
        while len(lines) < REQUIRED_LINES and retries < MAX_LINE_COUNT_RETRIES:
            retries += 1
            st.warning(
                f"⚠️ Блок {key}: получено {len(lines)}/{REQUIRED_LINES} строк. "
                f"Пере-генерирую (попытка {retries}/{MAX_LINE_COUNT_RETRIES})..."
            )
            fix_prompt, fix_missing, fix_expected, fix_json, fix_order = build_prompt(
                task["path"], base_dir, [key], niche_mode=niche_mode
            )
            if fix_missing:
                break
            fix_blocks = interact_with_ai_studio(fix_prompt, fix_expected, fix_json)
            if fix_blocks and len(fix_blocks) >= fix_expected:
                lines = clean_lines(fix_blocks[0])
                if len(lines) > REQUIRED_LINES:
                    lines = lines[:REQUIRED_LINES]

        if len(lines) < REQUIRED_LINES:
            st.error(
                f"❌ Не удалось получить {REQUIRED_LINES} строк для блока {key} "
                f"(осталось {len(lines)}) после {MAX_LINE_COUNT_RETRIES} попыток. Проверь вручную."
            )

        st.session_state[key_to_state[key]] = "\n".join(lines)

    st.session_state.just_regenerated.setdefault(task_index, set()).update(blocks_order)
    return True


# ─── МАССОВЫЙ АВТО-ПРОГОН БЕЗ ПАУЗ ───────────────────────────────────────────
if st.session_state.get("auto_run_selected", False):
    remain = sorted([i for i in st.session_state.selected_tasks if i not in st.session_state.done_set])
    if not remain:
        st.success("🎉 Массовый авто-прогон полностью завершен! Все Excel заполнены.")
        st.session_state.auto_run_selected = False
        st.balloons()
    else:
        cur_target = remain[0]
        st.session_state.cur_idx = cur_target
        st.toast(f"🤖 Авто-генерация [{cur_target + 1}/{len(tasks)}]: {tasks[cur_target]['strat']}")

        if run_ai_for_current_task(cur_target, "ALL"):
            t = clean_lines(st.session_state["ta_titles"])
            d = clean_lines(st.session_state["ta_descs"])
            c = clean_lines(st.session_state["ta_ctas"])
            if t and d and c:
                write_xlsx(tasks[cur_target]["xlsx"], t, d, c)
                st.session_state.done_set.add(cur_target)

                # Сбрасываем QC на wait и сохраняем
                st.session_state.qc_state[cur_target] = {"TITLES": "wait", "DESCS": "wait", "CTAS": "wait"}
                save_qc_state(tasks[cur_target]["path"], st.session_state.qc_state[cur_target])
            st.rerun()
        else:
            st.error(f"Сбой генерации для {tasks[cur_target]['strat']}. Авто-прогон остановлен.")
            st.session_state.auto_run_selected = False

# QC-АВТОПРОГОН ДЛЯ ИСПРАВЛЕНИЙ БЕЗ ПАУЗ
elif st.session_state.get("qc_run_redos", False):
    target_idx = None
    targets_to_redo = []
    for i in range(len(tasks)):
        q = st.session_state.qc_state[i]
        if "redo" in q.values():
            target_idx = i
            if q["TITLES"] == "redo": targets_to_redo.append("TITLES")
            if q["DESCS"] == "redo": targets_to_redo.append("DESCS")
            if q["CTAS"] == "redo": targets_to_redo.append("CTAS")
            break

    if target_idx is None:
        st.success("🎉 Все исправления по Контролю Качества завершены!")
        st.session_state.qc_run_redos = False
        st.balloons()
    else:
        st.session_state.cur_idx = target_idx
        update_text_areas(target_idx)

        st.toast(f"🔧 QC исправление {'+'.join(targets_to_redo)}: {tasks[target_idx]['strat']}")

        # Один вызов на ВСЕ помеченные блоки сразу (1, 2 или 3) — вместо цикла
        # с отдельным промптом на каждый блок.
        success = run_ai_for_current_task(target_idx, targets_to_redo)

        if success:
            for t_key in targets_to_redo:
                st.session_state.qc_state[target_idx][t_key] = "wait"
            save_qc_state(tasks[target_idx]["path"], st.session_state.qc_state[target_idx])

            t = clean_lines(st.session_state["ta_titles"])
            d = clean_lines(st.session_state["ta_descs"])
            c = clean_lines(st.session_state["ta_ctas"])
            write_xlsx(tasks[target_idx]["xlsx"], t, d, c)
            st.session_state.done_set.add(target_idx)
            st.rerun()
        else:
            st.session_state.qc_run_redos = False

# ─── SIDEBAR ──────────────────────────────────────────────────────────────────
with st.sidebar:
    st.markdown('<div class="main-header">🤖 Auto Filler</div>', unsafe_allow_html=True)
    if tasks: st.progress(len(st.session_state.done_set) / len(tasks),
                          text=f"Заполнено Excel: {len(st.session_state.done_set)} / {len(tasks)}")

    st.markdown("### 📋 Список папок")
    c_sel1, c_sel2 = st.columns(2)
    if c_sel1.button("✅ Выбрать все", use_container_width=True):
        st.session_state.selected_tasks = set(range(len(tasks)))
        st.rerun()
    if c_sel2.button("❌ Снять все", use_container_width=True):
        st.session_state.selected_tasks = set()
        st.rerun()

    st.caption("Галочка = папка участвует в Массовых действиях")

    last_niche = None
    for i, t in enumerate(tasks):
        if t["niche"] != last_niche:
            st.markdown(
                f"<div style='margin-top:14px;margin-bottom:4px;font-weight:bold;color:#a78bfa;font-size:13px;'>📁 {t['niche']}</div>",
                unsafe_allow_html=True)
            last_niche = t["niche"]

        col_cb, col_btn = st.columns([1, 6])
        is_sel = i in st.session_state.selected_tasks

        if col_cb.checkbox("Выбрать", value=is_sel, key=f"sel_{i}", label_visibility="collapsed") != is_sel:
            if is_sel:
                st.session_state.selected_tasks.remove(i)
            else:
                st.session_state.selected_tasks.add(i)
            st.rerun()

        icon = "✅" if i in st.session_state.done_set else ("📊" if t["xlsx"] else "❌")

        # Индикатор QC в сайдбаре (redo/wait перекрывают базовую иконку; wait важнее redo)
        q = st.session_state.qc_state[i]
        if "redo" in q.values(): icon = "🔴"
        if "wait" in q.values(): icon = "🕐"

        if i == cur:
            col_btn.markdown(
                f"<div style='background:#1e1030;border:1px solid #a78bfa;border-radius:6px;padding:4px;font-size:11px;color:#c4b5fd;'>▶ {t['strat']}</div>",
                unsafe_allow_html=True)
        else:
            if col_btn.button(f"{icon} {t['strat']}", key=f"nav_{i}", use_container_width=True):
                st.session_state.cur_idx = i
                update_text_areas(i)
                st.rerun()

# ─── MAIN APP (ГЕНЕРАЦИЯ + QC ВМЕСТЕ) ──────────────────────────────────────────
if not tasks: st.stop()
task = tasks[cur]

st.markdown("### ⚡ Массовые действия (без пауз)")
c1, c2 = st.columns(2)
with c1:
    if st.button("▶️ Заполнить ПУСТЫЕ выбранные (Автоматом)", type="primary", use_container_width=True):
        st.session_state.auto_run_selected = True
        st.rerun()
with c2:
    if st.button("⚠️ ПЕРЕЗАПИСАТЬ выбранные (Заново)", use_container_width=True):
        st.session_state.show_confirm_overwrite = True
        st.rerun()

if st.session_state.get("show_confirm_overwrite", False):
    st.warning(
        f"Начнется НЕПРЕРЫВНЫЙ авто-прогон для {len(st.session_state.selected_tasks)} папок. Старые данные в Excel будут перезаписаны.")
    cw1, cw2 = st.columns(2)
    if cw1.button("🔴 ДА, НАЧАТЬ МАССОВУЮ ПЕРЕЗАПИСЬ", type="primary"):
        st.session_state.show_confirm_overwrite = False
        st.session_state.done_set = st.session_state.done_set - st.session_state.selected_tasks
        st.session_state.auto_run_selected = True
        st.rerun()
    if cw2.button("Отмена"):
        st.session_state.show_confirm_overwrite = False
        st.rerun()

st.divider()

json_file_path = os.path.join(base_dir, "all_products.json")
if not os.path.exists(json_file_path): json_file_path = os.path.join(task['path'], "all_products.json")
json_text = f"📎 JSON: {os.path.basename(json_file_path)}" if os.path.exists(json_file_path) else "⚠️ JSON не найден!"

st.markdown(f"""
<div class="card card-accent">
  <div class="niche-badge">📁 {task['niche']} ➜ {task['strat']}</div>
  <div style="font-size: 11px; color: #34d399;">📊 {task['xlsx']}</div>
  <div style="font-size: 11px; color: #a78bfa; margin-top: 3px;">{json_text}</div>
</div>
""", unsafe_allow_html=True)

st.session_state.niche_mode = st.checkbox(
    "🧭 Бить в ОБЩУЮ НИШУ (авто-определение темы перед генерацией)",
    value=st.session_state.niche_mode,
    help="Включи, если обычная генерация промахивается мимо части фото из выдачи "
         "(например теги смешивают еду и напитки, а тексты получаются заточены только под 'десерты'). "
         "Gemini сначала сам определит общую широкую нишу по ВСЕМ тегам, покажет её (НИША: ...), "
         "и уже под неё сгенерирует тексты, подходящие к любой картинке по любому тегу из списка."
)

mc1, mc2 = st.columns(2)
with mc1:
    if st.button("🤖 Сгенерировать ВСЁ для этой папки", use_container_width=True):
        with st.spinner("Работаем с AI..."):
            if run_ai_for_current_task(cur, "ALL"): st.rerun()
with mc2:
    with st.popover("🔄 Точечная регенерация (исправить)"):
        if st.button("Переделать ЗАГОЛОВКИ"):
            if run_ai_for_current_task(cur, "TITLES"): st.rerun()
        if st.button("Переделать ОПИСАНИЯ"):
            if run_ai_for_current_task(cur, "DESCS"): st.rerun()
        if st.button("Переделать CTA"):
            if run_ai_for_current_task(cur, "CTAS"): st.rerun()

# --- ПОЛЯ И КОНТРОЛЬ КАЧЕСТВА ---
col1, col2, col3 = st.columns(3)


def qc_radio(label, current_val, key):
    opts = {"wait": "⏳ Ждет проверки", "ok": "🟢 ОК", "redo": "🔴 Переделать"}
    return st.radio(label, options=list(opts.keys()), format_func=lambda x: opts[x],
                    index=list(opts.keys()).index(current_val), key=key, horizontal=True)


def set_qc(t, d, c):
    st.session_state.qc_state[cur] = {"TITLES": t, "DESCS": d, "CTAS": c}
    save_qc_state(task["path"], st.session_state.qc_state[cur])


with col1:
    badge_t = " 🆕 ТОЛЬКО ЧТО ПЕРЕДЕЛАНО" if "TITLES" in st.session_state.just_regenerated.get(cur, set()) else ""
    raw_t = st.text_area(f"📝 ЗАГОЛОВКИ{badge_t}", height=340, key="ta_titles")
    titles = clean_lines(raw_t)
    st.caption(f"Строк: {len(titles)}")
    # QC Radio под полем
    new_t = qc_radio("Оценка:", st.session_state.qc_state[cur]["TITLES"], f"qct_{cur}")

with col2:
    badge_d = " 🆕 ТОЛЬКО ЧТО ПЕРЕДЕЛАНО" if "DESCS" in st.session_state.just_regenerated.get(cur, set()) else ""
    raw_d = st.text_area(f"📖 ОПИСАНИЯ{badge_d}", height=340, key="ta_descs")
    descs = clean_lines(raw_d)
    st.caption(f"Строк: {len(descs)}")
    # QC Radio под полем
    new_d = qc_radio("Оценка:", st.session_state.qc_state[cur]["DESCS"], f"qcd_{cur}")

with col3:
    badge_c = " 🆕 ТОЛЬКО ЧТО ПЕРЕДЕЛАНО" if "CTAS" in st.session_state.just_regenerated.get(cur, set()) else ""
    raw_c = st.text_area(f"🎯 CTA{badge_c}", height=340, key="ta_ctas")
    ctas = clean_lines(raw_c)
    st.caption(f"Строк: {len(ctas)}")
    # QC Radio под полем
    new_c = qc_radio("Оценка:", st.session_state.qc_state[cur]["CTAS"], f"qcc_{cur}")

# Сохранение изменения радиокнопок
if new_t != st.session_state.qc_state[cur]["TITLES"] or new_d != st.session_state.qc_state[cur]["DESCS"] or new_c != \
        st.session_state.qc_state[cur]["CTAS"]:
    if new_t != st.session_state.qc_state[cur]["TITLES"]:
        st.session_state.just_regenerated.get(cur, set()).discard("TITLES")
    if new_d != st.session_state.qc_state[cur]["DESCS"]:
        st.session_state.just_regenerated.get(cur, set()).discard("DESCS")
    if new_c != st.session_state.qc_state[cur]["CTAS"]:
        st.session_state.just_regenerated.get(cur, set()).discard("CTAS")
    set_qc(new_t, new_d, new_c)
    st.rerun()

st.divider()

# ─── ГЛОБАЛЬНЫЕ КНОПКИ QC И СОХРАНЕНИЯ ────────────────────────────────────────
qc_b1, qc_b2 = st.columns(2)
if qc_b1.button("✅ ПОМЕТИТЬ ВСЁ ИДЕАЛЬНЫМ", use_container_width=True):
    set_qc("ok", "ok", "ok")
    st.rerun()
if qc_b2.button("🔴 ОТПРАВИТЬ ВСЁ НА ПЕРЕДЕЛКУ", use_container_width=True):
    set_qc("redo", "redo", "redo")
    st.rerun()

def _dup_warning(label, lines):
    uniq = len(set(lines))
    if uniq < len(lines):
        st.warning(
            f"⚠️ В блоке {label} есть {len(lines) - uniq} дублирующихся строк(и) — "
            f"после сохранения при перезагрузке они схлопнутся, и уникальных останется {uniq}, а не {len(lines)}. "
            f"Если нужно ровно {REQUIRED_LINES} уникальных строк — замени дубликат на новый уникальный текст."
        )

_dup_warning("ЗАГОЛОВКИ", titles)
_dup_warning("ОПИСАНИЯ", descs)
_dup_warning("CTA", ctas)

write_ready = bool(titles and descs and ctas and task["xlsx"])
if st.button("💾 СОХРАНИТЬ ТЕКУЩУЮ СТРАНИЦУ В EXCEL", type="primary", use_container_width=True,
             disabled=not write_ready):
    write_xlsx(task["xlsx"], titles, descs, ctas)
    st.session_state.done_set.add(cur)
    # Оставляем QC таким, каким он был (или ставим OK, если хочешь)
    st.success("✅ Сохранено в Excel!")

st.divider()

# ─── БЛОК АВТОМАТИЧЕСКОГО ИСПРАВЛЕНИЯ БРАКА ─────────────────────────────────
total_redos = sum(1 for q in st.session_state.qc_state.values() for v in q.values() if v == "redo")

if total_redos > 0:
    st.markdown(
        f"**Блоков, отправленных на переделку (по всем папкам):** <span style='color:#f87171;font-weight:bold;'>{total_redos}</span>",
        unsafe_allow_html=True)
    if st.button("🚀 ЗАВЕРШИТЬ КОНТРОЛЬ И ИСПРАВИТЬ БРАК (Автоматом через ИИ)", type="primary",
                 use_container_width=True):
        st.session_state.show_confirm_qc = True
        st.rerun()

    if st.session_state.get("show_confirm_qc", False):
        st.warning(f"Начнется непрерывный авто-прогон для {total_redos} бракованных блоков.")
        c_qc1, c_qc2 = st.columns(2)
        if c_qc1.button("✅ ДА, ИСПРАВИТЬ ВСЁ"):
            st.session_state.show_confirm_qc = False
            st.session_state.qc_run_redos = True
            st.rerun()
        if c_qc2.button("❌ Отмена"):
            st.session_state.show_confirm_qc = False
            st.rerun()