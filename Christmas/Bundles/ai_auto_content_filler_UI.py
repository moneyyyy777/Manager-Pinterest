import sys
import subprocess
import os

# 1. ПРОВЕРЯЕМ АВТОЗАПУСК СТРОГО ДО ИМПОРТА STREAMLIT!
if "streamlit" not in sys.modules:
    script = os.path.abspath(__file__)
    subprocess.run([sys.executable, "-m", "streamlit", "run", script])
    sys.exit()

# 2. ИМПОРТИРУЕМ ВСЁ ОСТАЛЬНОЕ
import json
import time
import re
import openpyxl
import streamlit as st
import streamlit.components.v1 as components
from playwright.sync_api import sync_playwright, TimeoutError

# ─── НАСТРОЙКИ ────────────────────────────────────────────────────────────────
START_ROW = 2
END_ROW = 5001
COL_DESC = 12
COL_CTA = 13
COL_TITLE = 14
REQUIRED_LINES = 100

CHROME_DEBUG_URL = "http://127.0.0.1:9222"

st.set_page_config(page_title="🤖 Pinterest Filler & QC", layout="wide", initial_sidebar_state="expanded")

st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=JetBrains+Mono:wght@400;700&family=Syne:wght@700;800&display=swap');

/* Только шрифты! Жесткие цвета фона убраны, чтобы не ломать контраст в Светлой/Темной теме */
html, body, [class*="css"] { 
    font-family: 'JetBrains Mono', monospace !important; 
}

/* Красивая градиентная шапка */
.main-header { 
    font-family: 'Syne', sans-serif; 
    font-size: 26px; 
    font-weight: 800; 
    background: linear-gradient(90deg, #a78bfa, #f472b6); 
    -webkit-background-clip: text; 
    -webkit-text-fill-color: transparent; 
}

/* Плашка с папкой (всегда контрастная) */
.niche-badge {
    display: inline-block; 
    background: #5b21b6 !important;
    border: 1px solid #8b5cf6 !important; 
    border-radius: 6px;
    padding: 6px 14px; 
    font-size: 13px; 
    font-weight: 700;
    color: #ffffff !important; 
    margin-bottom: 8px;
}

/* Адаптивная карточка (полупрозрачная, отлично смотрится и на белом, и на черном) */
.card { 
    background-color: rgba(139, 92, 246, 0.05);
    border: 1px solid rgba(139, 92, 246, 0.2); 
    border-radius: 10px; 
    padding: 16px 20px; 
    margin-bottom: 12px; 
}
.card-accent { border-left: 3px solid #8b5cf6; }

/* Главная кнопка */
div.stButton > button[kind="primary"] { 
    background: linear-gradient(135deg, #7c3aed, #a855f7) !important; 
    border: none !important; 
    color: white !important; 
    font-weight: 700 !important; 
}

/* Убираем тусклость у заблокированных окон (теги и tasks.txt) */
textarea[disabled] {
    color: var(--text-color) !important;
    -webkit-text-fill-color: var(--text-color) !important;
    opacity: 1 !important;
}
</style>
""", unsafe_allow_html=True)


# ─── ИНТЕГРАЦИЯ С PLAYWRIGHT (AI STUDIO) ──────────────────────────────────────
def interact_with_ai_studio(prompt: str, expected_blocks: int, json_file_path: str, max_attempts: int = 15):
    with sync_playwright() as p:
        try:
            browser = p.chromium.connect_over_cdp(CHROME_DEBUG_URL)
            context = browser.contexts[0]

            page = None
            for p_ext in context.pages:
                if "aistudio" in p_ext.url:
                    page = p_ext
                    break
            if not page:
                page = context.new_page()

            for attempt in range(1, max_attempts + 1):
                try:
                    # Убрано page.bring_to_front() чтобы окно не дергалось и работало в фоне
                    st.toast(f"🔄 Попытка {attempt}/{max_attempts}: Подготовка AI Studio...")

                    page.goto("https://aistudio.google.com/app/u/1/prompts/new_chat", timeout=30000)
                    page.wait_for_timeout(2000)

                    if json_file_path and os.path.exists(json_file_path):
                        try:
                            file_input = page.locator('input[type="file"]').first
                            file_input.set_input_files(json_file_path)
                            page.wait_for_timeout(3000)
                        except Exception as e:
                            print(f"⚠️ Ошибка загрузки файла: {e}")

                    text_area = page.locator("textarea").last
                    text_area.fill(prompt)
                    page.wait_for_timeout(500)

                    try:
                        grounding_sw = page.locator('button[role="switch"][aria-label*="Grounding with Google Search"]')
                        if grounding_sw.is_visible(timeout=500) and grounding_sw.get_attribute(
                                "aria-checked") == "true":
                            grounding_sw.click()
                            page.wait_for_timeout(300)

                        url_sw = page.locator('button[role="switch"][aria-label*="Browse the url context"]')
                        if url_sw.is_visible(timeout=500) and url_sw.get_attribute("aria-checked") == "true":
                            url_sw.click()
                            page.wait_for_timeout(300)
                    except Exception:
                        pass

                    try:
                        clicked = page.evaluate('''() => {
                            const btn = document.querySelector('button[mattooltipclass="run-button-tooltip"]') 
                                     || document.querySelector('button.ctrl-enter-submits');
                            if (btn) { btn.click(); return true; }
                            return false;
                        }''')
                        if not clicked:
                            page.locator('button[mattooltipclass="run-button-tooltip"]').first.click(force=True)
                    except Exception:
                        text_area.focus()
                        page.keyboard.press("Meta+Enter")

                    st.toast(f"⏳ Попытка {attempt}: AI генерирует ответ...")

                    prev_len = 0
                    stable_ticks = 0

                    for _ in range(90):  # До 3 минут генерации
                        time.sleep(2)
                        page_content = page.content()
                        if "An internal error has occurred" in page_content or "Failed to generate content" in page_content:
                            raise Exception("AI Studio выдал внутреннюю ошибку.")

                        code_blocks = page.locator("ms-code-block pre").all_inner_texts()
                        if not code_blocks:
                            code_blocks = page.locator("pre").all_inner_texts()

                        if len(code_blocks) >= expected_blocks:
                            curr_len = len(code_blocks[-1])
                            if curr_len > 30 and curr_len == prev_len:
                                stable_ticks += 1
                                if stable_ticks >= 3:
                                    final_blocks = page.locator("ms-code-block pre").all_inner_texts()
                                    if not final_blocks:
                                        final_blocks = page.locator("pre").all_inner_texts()
                                    if len(final_blocks) != expected_blocks:
                                        raise Exception(
                                            f"AI вернул {len(final_blocks)} блоков кода вместо ожидаемых {expected_blocks} "
                                            f"— результат ненадёжен, переспрашиваю заново."
                                        )
                                    return final_blocks[-expected_blocks:]
                            else:
                                stable_ticks = 0
                                prev_len = curr_len

                    raise Exception("Таймаут ожидания ответа AI.")

                except Exception as attempt_err:
                    warning_placeholder = st.empty()
                    warning_placeholder.warning(f"⚠️ Попытка {attempt} сбой: {attempt_err}. Перезапуск через 5 сек...")
                    time.sleep(5)
                    warning_placeholder.empty()

            raise Exception(f"Все {max_attempts} попыток завершились ошибкой AI Studio.")
        except Exception as e:
            err_str = str(e)
            st.error("❌ Ошибка соединения с AI Studio или таймаут.")
            with st.expander("Скопировать полный лог ошибки (нажми тут)"):
                st.code(err_str, language="text")
            return []


# ─── ХЕЛПЕРЫ ──────────────────────────────────────────────────────────────────
def read_txt(path):
    if not os.path.exists(path): return ""
    with open(path, "r", encoding="utf-8", errors="replace") as f: return f.read().strip()

def clean_lines(raw):
    if not raw: return []
    lines = []
    for line in raw.splitlines():
        line = line.strip()
        line = re.sub(r'^\d+[\.\)]\s*|^[*\-•]\s*|^`{3,}', '', line).replace("```", "").strip()
        if line: lines.append(line)
    return lines

def load_qc_state(task_path):
    json_path = os.path.join(task_path, "qc_state.json")
    if os.path.exists(json_path):
        try:
            with open(json_path, "r", encoding="utf-8") as f:
                data = json.load(f)
                if "TAGS" not in data:
                    data["TAGS"] = "wait"
                return data
        except: pass
    return {"TITLES": "wait", "DESCS": "wait", "CTAS": "wait", "TAGS": "wait"}

def save_qc_state(task_path, state):
    json_path = os.path.join(task_path, "qc_state.json")
    try:
        with open(json_path, "w", encoding="utf-8") as f:
            json.dump(state, f, ensure_ascii=False, indent=2)
    except Exception as e:
        print(f"Ошибка сохранения QC: {e}")

def scan_tasks(base_dir):
    tasks = []
    if not os.path.isdir(base_dir): return tasks
    niches = sorted(
        d for d in os.listdir(base_dir) if os.path.isdir(os.path.join(base_dir, d)) and not d.startswith("."))
    for niche in niches:
        np = os.path.join(base_dir, niche)
        try:
            strats = sorted(d for d in os.listdir(np) if os.path.isdir(os.path.join(np, d)) and re.match(r"^\d+", d))
        except PermissionError: continue
        for s in strats:
            sp = os.path.join(np, s)
            xlsx = next((os.path.join(sp, f) for f in os.listdir(sp) if f.endswith(".xlsx") and not f.startswith("~$")), None)
            tasks.append({"niche": niche, "strat": s, "path": sp, "xlsx": xlsx})
    return tasks

def check_xlsx_filled(xlsx_path):
    if not xlsx_path: return False
    try:
        wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
        ws = wb.active
        filled = sum(1 for row in range(START_ROW, min(START_ROW + 3, END_ROW + 1))
                     if ws.cell(row=row, column=COL_TITLE).value and ws.cell(row=row, column=COL_DESC).value and ws.cell(row=row, column=COL_CTA).value)
        wb.close()
        return filled >= 3
    except:
        return False

def read_unique_xlsx_data(xlsx_path):
    if not xlsx_path or not os.path.exists(xlsx_path): return "", "", ""
    titles, descs, ctas = [], [], []
    seen_t, seen_d, seen_c = set(), set(), set()
    try:
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
        ws = wb.active
        max_search_row = min(START_ROW + 500, ws.max_row + 1)
        for row in range(START_ROW, max_search_row):
            d_val = ws.cell(row=row, column=COL_DESC).value
            c_val = ws.cell(row=row, column=COL_CTA).value
            t_val = ws.cell(row=row, column=COL_TITLE).value
            if t_val and str(t_val).strip() not in seen_t:
                titles.append(str(t_val).strip())
                seen_t.add(str(t_val).strip())
            if d_val and str(d_val).strip() not in seen_d:
                descs.append(str(d_val).strip())
                seen_d.add(str(d_val).strip())
            if c_val and str(c_val).strip() not in seen_c:
                ctas.append(str(c_val).strip())
                seen_c.add(str(c_val).strip())
        wb.close()
        return "\n".join(titles), "\n".join(descs), "\n".join(ctas)
    except:
        return "", "", ""

def update_text_areas(task_idx):
    task = st.session_state.tasks[task_idx]
    t, d, c = read_unique_xlsx_data(task["xlsx"])
    st.session_state["ta_titles"] = t
    st.session_state["ta_descs"] = d
    st.session_state["ta_ctas"] = c
    if "render_keys" in st.session_state:
        st.session_state.render_keys["ta_titles"] += 1
        st.session_state.render_keys["ta_descs"] += 1
        st.session_state.render_keys["ta_ctas"] += 1

def _auto_detect_context(strat_path, base_dir):
    """Автоматически определяет: тип файла (Filter/Intent), стоп-слова из соседних подниш, тип продукта из тегов."""
    strat_name = os.path.basename(strat_path)
    niche_dir = os.path.dirname(strat_path)  # Папка подниши (напр. Bundles Christmas Planners)
    parent_dir = os.path.dirname(niche_dir)    # Папка уровня выше (напр. Bundles)

    # --- ТИП ФАЙЛА ---
    is_filter = "Filter" in strat_name or "filter" in strat_name
    file_type = "ОБЩИЙ (Filter by OpenAI)" if is_filter else "ИНТЕНТ"

    # --- СТОП-СЛОВА из СОСЕДНИХ ПОДНИШ ---
    stop_words = set()
    niche_basename = os.path.basename(niche_dir)
    if os.path.isdir(parent_dir):
        for sibling in os.listdir(parent_dir):
            sibling_path = os.path.join(parent_dir, sibling)
            if not os.path.isdir(sibling_path) or sibling == niche_basename:
                continue
            # Извлекаем ключевые слова из имени соседней подниши
            # Убираем общий префикс (напр. "Bundles Christmas") и берем уникальную часть
            parts = sibling.split()
            # Берем слова, которых нет в названии текущей подниши
            for word in parts:
                w_lower = word.lower()
                if len(w_lower) > 2 and w_lower not in niche_basename.lower():
                    stop_words.add(w_lower)

    # Всегда добавляем стандартные стоп-слова (маркетплейсы, ПО, годы)
    always_stop = {
        "etsy", "creative fabrica", "creative market", "design bundles", "font bundles",
        "canva", "adobe", "photoshop", "illustrator", "goodnotes", "notion",
        "2021", "2022", "2023", "2024", "2025"
    }
    stop_words.update(always_stop)

    # --- ТИП ПРОДУКТА из тегов ---
    for_main = os.path.join(strat_path, "For Main")
    tags_text = read_txt(os.path.join(for_main, "Boards_Tags.txt")).lower()
    product_types = []
    product_keywords = {
        "svg": ["svg", "svgs", "cut file", "cut files"],
        "clipart": ["clipart", "clip art"],
        "font": ["font", "fonts", "typeface", "typography"],
        "planner": ["planner", "planners"],
        "sublimation": ["sublimation", "sublimate"],
        "sticker": ["sticker", "stickers"],
        "journal": ["journal", "journaling"],
        "pattern": ["pattern", "patterns", "seamless"],
    }
    for product, keywords in product_keywords.items():
        if any(kw in tags_text for kw in keywords):
            product_types.append(product)

    return file_type, sorted(stop_words), product_types


def build_prompt(strat_path, base_dir, target="ALL", niche_mode=False, excluded_niches="", use_json=True, niche_hint=""):
    for_main = os.path.join(strat_path, "For Main")
    link = read_txt(os.path.join(base_dir, "linksite.txt"))
    tags = read_txt(os.path.join(for_main, "Boards_Tags.txt"))
    board = read_txt(os.path.join(for_main, "Boards_Name.txt"))

    missing = []
    if not link: missing.append("linksite.txt")
    if not tags: missing.append("Boards_Tags.txt")
    if not board: missing.append("Boards_Name.txt")

    json_file = os.path.join(base_dir, "all_products.json")
    if not os.path.exists(json_file):
        json_file = os.path.join(strat_path, "all_products.json")
        if not os.path.exists(json_file): json_file = None
    if not use_json:
        json_file = None

    if missing: return "", missing, 0, None, []

    # ═══ АВТОКОНТЕКСТ: тип файла, стоп-слова, тип продукта ═══
    file_type, auto_stop_words, product_types = _auto_detect_context(strat_path, base_dir)
    product_label = ", ".join(product_types).title() if product_types else "Digital Design Resources"

    TASK_LABELS = {
        "TITLES": (
            "РОВНО 100 ЗАГОЛОВКОВ (не 99, не 80, не 105 — точно 100 строк)",
            "без эмодзи, без знаков препинания в конце, 3-8 слов, содержат ключевые слова из тегов естественно вписанные"
        ),
        "DESCS": (
            "РОВНО 100 ОПИСАНИЙ (не 99, не 80, не 105 — точно 100 строк)",
            "⚠️ ОБЯЗАТЕЛЬНО в КАЖДОЙ строке ставь эмодзи по смыслу к содержанию и тегам в самом начале предложения — "
            "это требование не снимается ни при каких условиях, без артикля \"A\" в начале, без призыва к действию внутри, "
            "1-2 предложения, должны подходить к ЛЮБОЙ картинке из выдачи по данным тегам."
        ),
        "CTAS": (
            "РОВНО 100 ПРИЗЫВОВ К ДЕЙСТВИЮ — CTA (не 99, не 80, не 105 — точно 100 строк)",
            "⚠️ ОБЯЗАТЕЛЬНО в КАЖДОЙ строке ставь эмодзи по смыслу в самом начале — "
            "это требование не снимается ни при каких условиях, одно короткое предложение, призыв перейти скачать наш товар "
            "бесплатно на сайте с подборкой лучшего товара, разные формулировки — не повторяться. "
            f"Тип товара с Заглавной (например: {product_label})"
        ),
    }
    CANONICAL_ORDER = ["TITLES", "DESCS", "CTAS"]

    if target == "ALL": blocks_order = CANONICAL_ORDER[:]
    elif isinstance(target, str): blocks_order = [target]
    else: blocks_order = [k for k in CANONICAL_ORDER if k in target]

    tasks_text = ""
    for idx, key in enumerate(blocks_order, start=1):
        name, desc = TASK_LABELS[key]
        tasks_text += f"{idx}) {name} ({desc})\n"
    expected = len(blocks_order)

    # ═══ БЛОК: Подсказка по нише ═══
    niche_hint_line = ""
    if niche_mode and niche_hint and niche_hint.strip():
        niche_hint_line = (
            f"💡 ПОДСКАЗКА ПО НИШЕ: если по смыслу тегов ниша может трактоваться неоднозначно — "
            f"предпочти формулировку ближе к '{niche_hint.strip()}', ЕСЛИ она подтверждается тегами.\n\n"
        )

    # ═══ БЛОК: Исключения ═══
    excluded_block = ""
    # Объединяем ручные исключения с автоматическими стоп-словами
    all_excluded = set()
    if excluded_niches and excluded_niches.strip():
        for term in excluded_niches.strip().split(","):
            all_excluded.add(term.strip())
    all_excluded.update(auto_stop_words)

    if all_excluded:
        excluded_block = (
            f"\n🚫 СТОП-СЛОВА И ЗАПРЕЩЁННЫЕ ТЕМЫ (строго нельзя упоминать в текстах):\n"
            f"{', '.join(sorted(all_excluded))}\n"
            f"Эти слова относятся к СОСЕДНИМ поднишам, которые продвигаются ОТДЕЛЬНО, к конкурирующему ПО или к устаревшим годам. "
            f"НИКОГДА не используй их в текстах. Если тема граничит с одним из стоп-слов — используй более общий термин.\n\n"
        )

    # ═══ БЛОК: Определение ниши ═══
    niche_block = ""
    if niche_mode:
        niche_block = (
            f"ШАГ 0 — ОПРЕДЕЛИ ОБЩУЮ НИШУ (сделай это ПЕРЕД генерацией текстов):\n"
            f"Внимательно прочитай ВЕСЬ список тегов целиком. "
            f"Определи САМУЮ ШИРОКУЮ тему, которая описывает ВСЕ теги из списка.\n"
            f"{niche_hint_line}"
            f"Напиши ПЕРВОЙ строкой ответа: НИША: <общая ниша на английском>\n\n"
            f"Дальше генерируй тексты под эту широкую нишу. Каждый текст должен подходить к ЛЮБОЙ картинке "
            f"по ЛЮБОМУ тегу из списка одновременно.\n\n"
        )

    # ═══ БЛОК: Тип файла и глубина ═══
    depth_block = ""
    if file_type == "ОБЩИЙ (Filter by OpenAI)":
        depth_block = (
            f"📌 ТИП ФАЙЛА: ОБЩИЙ (Filter by OpenAI) — теги РАЗНОРОДНЫЕ.\n"
            f"Пиши ШИРЕ. Тексты должны работать для ВСЕХ тегов одновременно. "
            f"Используй слова уровня ниши, не уходи в узкую под-тему.\n\n"
        )
    else:
        depth_block = (
            f"📌 ТИП ФАЙЛА: ИНТЕНТ — теги ОДНОРОДНЫЕ.\n"
            f"Можно писать ЧУТЬ конкретнее, но не привязывайся к одному объекту. "
            f"Картинки внутри интента похожи, но всё равно разные.\n\n"
        )

    json_line = (
        "Я прикрепил файл json с информацией по нашей подборке товара на сайте "
        "(детально изучи его, чтобы понять что пользователь получит при переходе с Pinterest).\n"
        if (use_json and json_file) else ""
    )

    prompt = (
        f"Ты — Мастер-Оркестратор для Pinterest. Я продвигаю товар через массовые аккаунты.\n"
        f"{json_line}\n"
        f"═══ КОНТЕКСТ БИЗНЕСА ═══\n"
        f"Мы скрапим картинки из Pinterest по тегам, размещаем их на аккаунтах с Title + Description + CTA. "
        f"CTA содержит ссылку на наш сайт. Мы НЕ знаем заранее, какая картинка попадётся — "
        f"поэтому тексты должны подходить к ЛЮБОЙ картинке из выдачи по тегам.\n\n"
        f"📋 Доска: {board}\n"
        f"🔗 Ссылка на сайт: {link}\n"
        f"{depth_block}"
        f"Теги доски (контекст для ключевых слов, НЕ вставляй дословно в тексты):\n"
        f"{tags}\n"
        f"{excluded_block}"
        f"{niche_block}"
        f"═══ ЗАДАНИЕ ═══\n"
        f"Сделай на английском:\n"
        f"{tasks_text}\n"
        f"═══ ПРАВИЛА ФОРМАТА ═══\n"
        f"- Каждый список помести в ОТДЕЛЬНЫЙ блок кода (```). В ответе РОВНО {expected} блок(ов) кода.\n"
        f"- Без нумерации, без пустых строк между строками внутри блока.\n"
        f"- Только сами тексты, ничего лишнего.\n\n"
        f"═══ ПРАВИЛА КАЧЕСТВА (КРИТИЧЕСКИ ВАЖНО) ═══\n"
        f"1. УНИКАЛЬНОСТЬ: ВСЕ 100 строк в каждом блоке уникальны. Нельзя написать одну фразу "
        f"и повторить с разными эмодзи. Используй разные синонимы, разную структуру предложений, "
        f"разные углы (качество, стиль, применение, эмоция, выгода, тренд).\n\n"
        f"2. ЗОЛОТОЕ ПРАВИЛО: Рандомный title[17] + рандомный desc[42] + рандомный cta[88] "
        f"должны ИДЕАЛЬНО сочетаться и подходить к ЛЮБОЙ картинке из выдачи.\n\n"
        f"3. ПРИНЦИП МАКСИМАЛЬНОЙ ГЛУБИНЫ: Погружайся в тематику ГЛУБОКО, но останавливайся "
        f"в точке, где ЛЮБОЙ пин из ЛЮБОГО тега всё ещё подходит под текст. "
        f"Не слишком общо (скучно) и не слишком узко (не подойдёт к большинству картинок).\n\n"
        f"4. ВИЗУАЛЬНЫЙ СТИЛЬ: Слова типа «watercolor», «boho», «vintage», «minimalist» — "
        f"используй ТОЛЬКО если они есть в ПОДАВЛЯЮЩЕМ БОЛЬШИНСТВЕ тегов. Иначе пиши нейтрально: "
        f"«digital», «beautiful», «stunning designs».\n\n"
        f"5. ДАТЫ: Мы живем в 2026-2027 годах. Категорически НЕ используй 2021-2025!\n\n"
        f"6. В CTA: В 70% строк используй слово FREE (БОЛЬШИМИ БУКВАМИ). "
        f"Название типа товара с Заглавной ({product_label}). "
        f"НЕ упоминай названия ПО (Canva, Adobe, Procreate, GoodNotes).\n\n"
        f"7. ЭМОДЗИ: В Описаниях и CTA — эмодзи в начале КАЖДОЙ строки. ⚠️ ПРАВИЛО РЕЛЕВАНТНОСТИ И РАЗНООБРАЗИЯ: Эмодзи должны СТРОГО соответствовать тематике ниши! Ищи максимально широкое разнообразие внутри логики ниши (уникально и красиво), но без абсурда. Недопустимо лениво крутить 3-4 одинаковых эмодзи на все 100 строк — используй весь арсенал релевантных (например, 15-25 идеально подходящих). Для CTA также разрешается миксовать их со стандартными «кнопочными» эмодзи конверсии для Pinterest которые дают большую конверсию.\n"
        f"В Заголовках — БЕЗ эмодзи, БЕЗ знаков препинания в конце.\n\n"
        f"═══ САМОПРОВЕРКА ПЕРЕД ОТПРАВКОЙ ═══\n"
        f"Перед выводом ОБЯЗАТЕЛЬНО проверь:\n"
        f"✅ В каждом блоке РОВНО 100 строк (не 80, не 99, не 105)? Если нет — допиши/удали.\n"
        f"✅ Все 100 строк уникальны (нет дублей даже с разными эмодзи)?\n"
        f"✅ В desc и CTA эмодзи стоит в начале КАЖДОЙ строки без исключения?\n"
        f"✅ Ни одно стоп-слово не просочилось?\n"
        f"✅ Нет старых годов (2021-2025)?\n"
        f"Если хоть одна проверка FAIL — исправь перед выводом.\n"
    )
    return prompt, missing, expected, json_file, blocks_order

def write_xlsx(xlsx_path, titles, descs, ctas):
    if not titles or not descs or not ctas:
        raise ValueError("Одно из полей пустое! Невозможно записать в Excel.")
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active
    total = END_ROW - START_ROW + 1
    for i in range(total):
        row = START_ROW + i
        ws.cell(row=row, column=COL_DESC).value = descs[i % len(descs)]
        ws.cell(row=row, column=COL_CTA).value = ctas[i % len(ctas)]
        ws.cell(row=row, column=COL_TITLE).value = titles[i % len(titles)]
    wb.save(xlsx_path)

# ─── ИНИЦИАЛИЗАЦИЯ СОСТОЯНИЯ ──────────────────────────────────────────────────
base_dir = os.path.dirname(os.path.abspath(__file__))

if "tasks" not in st.session_state:
    st.session_state.tasks = scan_tasks(base_dir)

if "qc_state" not in st.session_state:
    st.session_state.qc_state = {}
    for i, t in enumerate(st.session_state.tasks):
        st.session_state.qc_state[i] = load_qc_state(t["path"])
        if t["xlsx"]:
            tt, dd, cc = read_unique_xlsx_data(t["xlsx"])
            counts = {
                "TITLES": len(clean_lines(tt)),
                "DESCS": len(clean_lines(dd)),
                "CTAS": len(clean_lines(cc)),
            }
            changed = False
            for key, cnt in counts.items():
                if cnt > 0 and cnt != REQUIRED_LINES and st.session_state.qc_state[i][key] != "redo":
                    st.session_state.qc_state[i][key] = "redo"
                    changed = True
            if changed:
                save_qc_state(t["path"], st.session_state.qc_state[i])

if "done_set" not in st.session_state:
    st.session_state.done_set = set()
    for i, t in enumerate(st.session_state.tasks):
        if check_xlsx_filled(t["xlsx"]):
            qc = st.session_state.qc_state[i]
            # Если папка отправлена на полную перезапись, не добавляем ее в done_set
            if qc.get("TITLES") == "redo" and qc.get("DESCS") == "redo" and qc.get("CTAS") == "redo":
                pass
            else:
                st.session_state.done_set.add(i)

if "cur_idx" not in st.session_state:
    remain = sorted(set(range(len(st.session_state.tasks))) - st.session_state.done_set)
    st.session_state.cur_idx = remain[0] if remain else 0

if "selected_tasks" not in st.session_state:
    st.session_state.selected_tasks = {i for i in range(len(st.session_state.tasks)) if
                                       i not in st.session_state.done_set}

if "ta_titles" not in st.session_state:
    update_text_areas(st.session_state.cur_idx)

if "render_keys" not in st.session_state:
    st.session_state.render_keys = {"ta_titles": 0, "ta_descs": 0, "ta_ctas": 0}

tasks = st.session_state.tasks
cur = st.session_state.cur_idx

def get_niche_hint_for_task(task_index):
    terms = []
    seen_lower = set()
    for rule in st.session_state.niche_hint_rules:
        applies = rule["scope"] == "all" or task_index in rule["scope"]
        if not applies: continue
        for part in rule["text"].split(","):
            part = part.strip()
            if part and part.lower() not in seen_lower:
                seen_lower.add(part.lower())
                terms.append(part)
    return ", ".join(terms)

def get_excluded_niches_for_task(task_index):
    terms = []
    seen_lower = set()
    for rule in st.session_state.niche_exclusion_rules:
        applies = rule["scope"] == "all" or task_index in rule["scope"]
        if not applies: continue
        for part in rule["text"].split(","):
            part = part.strip()
            if part and part.lower() not in seen_lower:
                seen_lower.add(part.lower())
                terms.append(part)
    return ", ".join(terms)

if "niche_mode" not in st.session_state: st.session_state.niche_mode = False
if "niche_exclusion_rules" not in st.session_state: st.session_state.niche_exclusion_rules = []
if "niche_hint_rules" not in st.session_state: st.session_state.niche_hint_rules = []
if "use_json" not in st.session_state: st.session_state.use_json = True
if "just_regenerated" not in st.session_state: st.session_state.just_regenerated = {}

MAX_LINE_COUNT_RETRIES = 3

# ─── ЛОГИКА PLAYWRIGHT ДЛЯ UI ─────────────────────────────────────────────────
def run_ai_for_current_task(task_index, target="ALL", niche_mode=None):
    task = tasks[task_index]
    if niche_mode is None:
        niche_mode = st.session_state.get("niche_mode", False)
    excluded_niches = get_excluded_niches_for_task(task_index)
    niche_hint = get_niche_hint_for_task(task_index)
    use_json = st.session_state.get("use_json", True)
    prompt, missing, expected_blocks, json_file, blocks_order = build_prompt(
        task["path"], base_dir, target, niche_mode=niche_mode, excluded_niches=excluded_niches, use_json=use_json,
        niche_hint=niche_hint
    )
    if missing:
        st.session_state["_last_ai_error"] = f"❌ Не хватает файлов: {', '.join(missing)}"
        return False

    blocks = interact_with_ai_studio(prompt, expected_blocks, json_file)
    if not blocks or len(blocks) < expected_blocks:
        st.session_state["_last_ai_error"] = (
            f"❌ AI Studio не вернул нужное количество блоков кода для {task['strat']}. "
            f"Ожидалось: {expected_blocks}, получено: {len(blocks) if blocks else 0}. "
            f"Попробуй ещё раз — возможно, модель выдала текст без обёртки в блок кода."
        )
        return False

    key_to_state = {"TITLES": "ta_titles", "DESCS": "ta_descs", "CTAS": "ta_ctas"}

    for i, key in enumerate(blocks_order):
        lines = clean_lines(blocks[i])
        if len(lines) > REQUIRED_LINES:
            lines = lines[:REQUIRED_LINES]

        retries = 0
        while len(lines) < REQUIRED_LINES and retries < MAX_LINE_COUNT_RETRIES:
            retries += 1
            line_warn = st.empty()
            line_warn.warning(f"⚠️ Блок {key}: получено {len(lines)}/{REQUIRED_LINES} строк. Пере-генерирую (попытка {retries}/{MAX_LINE_COUNT_RETRIES})...")
            fix_prompt, fix_missing, fix_expected, fix_json, fix_order = build_prompt(
                task["path"], base_dir, [key], niche_mode=niche_mode, niche_hint=niche_hint
            )
            if fix_missing:
                line_warn.empty()
                break
            fix_blocks = interact_with_ai_studio(fix_prompt, fix_expected, fix_json)
            line_warn.empty()
            if fix_blocks and len(fix_blocks) >= fix_expected:
                lines = clean_lines(fix_blocks[0])
                if len(lines) > REQUIRED_LINES:
                    lines = lines[:REQUIRED_LINES]

        if len(lines) < REQUIRED_LINES:
            st.session_state["_last_ai_error"] = (
                f"⚠️ Блок {key}: получено {len(lines)}/{REQUIRED_LINES} строк после {MAX_LINE_COUNT_RETRIES} попыток. "
                f"Данные записаны как есть — проверь и при необходимости переделай этот блок."
            )

        # Записываем обновленные (и обрезанные до 100) строки в session_state
        state_key = key_to_state[key]
        st.session_state[state_key] = "\n".join(lines)
        if "render_keys" in st.session_state:
            st.session_state.render_keys[state_key] += 1

    st.session_state.just_regenerated.setdefault(task_index, set()).update(blocks_order)
    
    # ═══ FIX: Автоматический сброс статуса QC в "Ждет проверки" ═══
    for key in blocks_order:
        if task_index in st.session_state.qc_state:
            st.session_state.qc_state[task_index][key] = "wait"
    save_qc_state(task["path"], st.session_state.qc_state[task_index])

    # Очищаем прошлую ошибку при успехе
    st.session_state.pop("_last_ai_error", None)
    return True


# ─── МАССОВЫЙ АВТО-ПРОГОН БЕЗ ПАУЗ ───────────────────────────────────────────
if "success_log" not in st.session_state:
    st.session_state.success_log = []

if st.session_state.success_log:
    for msg in st.session_state.success_log[-5:]:  # Показываем последние 5 успешных действий
        st.success(msg)

if st.session_state.get("auto_run_selected", False):
    remain = sorted([i for i in st.session_state.selected_tasks if i not in st.session_state.done_set])
    if not remain:
        st.success("🎉 Массовый авто-прогон полностью завершен! Все Excel заполнены.")
        st.session_state.auto_run_selected = False
        st.session_state.success_log.clear()
        st.balloons()
    else:
        cur_target = remain[0]
        st.session_state.cur_idx = cur_target
        st.toast(f"🤖 Авто-генерация [{cur_target + 1}/{len(tasks)}]: {tasks[cur_target]['strat']}")
        if run_ai_for_current_task(cur_target, "ALL"):
            t = clean_lines(st.session_state["ta_titles"])
            d = clean_lines(st.session_state["ta_descs"])
            c = clean_lines(st.session_state["ta_ctas"])
            if t and d and c:
                write_xlsx(tasks[cur_target]["xlsx"], t, d, c)
                st.session_state.done_set.add(cur_target)
                old_tag = st.session_state.qc_state[cur_target].get("TAGS", "wait")
                st.session_state.qc_state[cur_target] = {"TITLES": "wait", "DESCS": "wait", "CTAS": "wait", "TAGS": old_tag}
                save_qc_state(tasks[cur_target]["path"], st.session_state.qc_state[cur_target])
            st.rerun()
        else:
            st.error(f"Сбой генерации для {tasks[cur_target]['strat']}. Авто-прогон остановлен.")
            st.session_state.auto_run_selected = False

# QC-АВТОПРОГОН ДЛЯ ИСПРАВЛЕНИЙ БЕЗ ПАУЗ
elif st.session_state.get("qc_run_redos", False):
    target_idx = None
    targets_to_redo = []
    for i in range(len(tasks)):
        q = st.session_state.qc_state[i]
        e_vals = [q.get("TITLES", ""), q.get("DESCS", ""), q.get("CTAS", "")]
        if "redo" in e_vals:
            target_idx = i
            if q.get("TITLES") == "redo": targets_to_redo.append("TITLES")
            if q.get("DESCS") == "redo": targets_to_redo.append("DESCS")
            if q.get("CTAS") == "redo": targets_to_redo.append("CTAS")
            break
    if target_idx is None:
        st.success("🎉 Все исправления по Контролю Качества завершены!")
        st.session_state.qc_run_redos = False
        st.balloons()
    else:
        st.session_state.cur_idx = target_idx
        update_text_areas(target_idx)
        st.toast(f"🔧 QC исправление {'+'.join(targets_to_redo)}: {tasks[target_idx]['strat']}")
        success = run_ai_for_current_task(target_idx, targets_to_redo)
        if success:
            for t_key in targets_to_redo: st.session_state.qc_state[target_idx][t_key] = "wait"
            save_qc_state(tasks[target_idx]["path"], st.session_state.qc_state[target_idx])
            t = clean_lines(st.session_state["ta_titles"])
            d = clean_lines(st.session_state["ta_descs"])
            c = clean_lines(st.session_state["ta_ctas"])
            write_xlsx(tasks[target_idx]["xlsx"], t, d, c)
            st.session_state.done_set.add(target_idx)
            st.rerun()
        else:
            st.session_state.qc_run_redos = False

# ─── SIDEBAR ──────────────────────────────────────────────────────────────────
with st.sidebar:
    st.markdown('<div class="main-header">🤖 Auto Filler</div>', unsafe_allow_html=True)
    if tasks:
        total = len(tasks)
        filled_excel = len(st.session_state.done_set)
        checked_excel = sum(1 for i in range(total) if all(st.session_state.qc_state[i][k] == "ok" for k in ["TITLES", "DESCS", "CTAS"]))
        checked_tags = sum(1 for i in range(total) if st.session_state.qc_state[i].get("TAGS") == "ok")
        
        st.progress(filled_excel / total, text=f"📝 Заполнено Excel: {filled_excel} / {total}")
        st.progress(checked_excel / total, text=f"✅ Проверено Excel: {checked_excel} / {total}")
        st.progress(checked_tags / total, text=f"🏷️ Проверено Тегов: {checked_tags} / {total}")

    st.divider()
    st.markdown("### 📋 Список папок")
    c_sel1, c_sel2 = st.columns(2)
    if c_sel1.button("✅ Выбрать все", use_container_width=True):
        st.session_state.selected_tasks = set(range(len(tasks)))
        st.rerun()
    if c_sel2.button("❌ Снять все", use_container_width=True):
        st.session_state.selected_tasks = set()
        st.rerun()
    st.caption("Галочка = папка участвует в Массовых действиях")

    last_niche = None
    for i, t in enumerate(tasks):
        if t["niche"] != last_niche:
            st.markdown(f"<div style='margin-top:14px;margin-bottom:4px;font-weight:bold;color:#8b5cf6;font-size:13px;'>📁 {t['niche']}</div>", unsafe_allow_html=True)
            last_niche = t["niche"]

        col_cb, col_btn = st.columns([1, 6])
        is_sel = i in st.session_state.selected_tasks

        if col_cb.checkbox("Выбрать", value=is_sel, key=f"sel_{i}", label_visibility="collapsed") != is_sel:
            if is_sel: st.session_state.selected_tasks.remove(i)
            else: st.session_state.selected_tasks.add(i)
            st.rerun()

        q = st.session_state.qc_state[i]
        
        # Статус тегов
        t_state = q.get("TAGS", "wait")
        if t_state == "ok": t_icon = "✅"
        elif t_state == "redo": t_icon = "🔴"
        else: t_icon = "🕐"
        
        # Статус Excel
        e_vals = [q["TITLES"], q["DESCS"], q["CTAS"]]
        if "redo" in e_vals: e_icon = "🔴"
        elif "wait" in e_vals: e_icon = "🕐"
        else: e_icon = "✅" if i in st.session_state.done_set else "➖"

        btn_label = f"🏷️{t_icon} 📊{e_icon} | {t['strat']}"

        if i == cur:
            col_btn.markdown(f"<div style='background:rgba(139, 92, 246, 0.15);border:1px solid #8b5cf6;border-radius:6px;padding:6px;font-size:12px;color:inherit;font-weight:bold;display:flex;align-items:center;'>▶ {btn_label}</div>", unsafe_allow_html=True)
        else:
            if col_btn.button(btn_label, key=f"nav_{i}", use_container_width=True):
                st.session_state.cur_idx = i
                update_text_areas(i)
                st.rerun()

    # ─── МЕНЕДЖЕР tasks.txt ДЛЯ ОРКЕСТРАТОРА (САЙДБАР) ────────────────────────
    st.divider()
    st.markdown("### 🚀 Оркестратор (tasks.txt)")
    
    WORKSPACE_ROOT = "/Users/kalifornia/Desktop/Pinterest/Managers Pinterest"
    tasks_txt_path = os.path.join(WORKSPACE_ROOT, "Websites/CloudFlare/artify-studio/scripts/artify_sites_publishing/tasks.txt")
    
    def read_tasks_txt():
        if not os.path.exists(tasks_txt_path): return []
        with open(tasks_txt_path, "r", encoding="utf-8") as f:
            return [line.strip() for line in f if line.strip()]
            
    def write_tasks_txt(lines):
        # Создаем папку, если вдруг её нет
        os.makedirs(os.path.dirname(tasks_txt_path), exist_ok=True)
        with open(tasks_txt_path, "w", encoding="utf-8") as f:
            f.write("\n".join(lines) + "\n")

    current_tasks_txt = read_tasks_txt()
    
    # 📌 Формируем пути относительно корня воркспейса (чтобы Оркестратор их правильно понял)
    all_options = []
    for t in tasks:
        full_path = os.path.join(base_dir, t["niche"])
        rel_path = os.path.relpath(full_path, WORKSPACE_ROOT)
        if rel_path not in all_options:
            all_options.append(rel_path)
    all_options = sorted(list(set(all_options)))

    default_selection = [opt for opt in all_options if opt in current_tasks_txt]
    if not current_tasks_txt and not os.path.exists(tasks_txt_path):
        default_selection = all_options 
        
    selected_for_orch = st.multiselect(
        "Пути для запуска (Только родители):",
        options=all_options,
        default=default_selection,
        help="Выбери папки, которые Оркестратор возьмет в работу."
    )
    
    btn_col1, btn_col2 = st.columns([2, 1])
    if btn_col1.button("💾 Записать", type="primary", use_container_width=True):
        new_tasks = []
        for line in current_tasks_txt:
            if line not in all_options: new_tasks.append(line)
            elif line in selected_for_orch: new_tasks.append(line)
        for s in selected_for_orch:
            if s not in new_tasks: new_tasks.append(s)
        write_tasks_txt(new_tasks)
        st.success("✅ Сохранено в Оркестратор!")
        st.rerun()
        
    if btn_col2.button("🗑️ Очистить", use_container_width=True):
        write_tasks_txt([])
        st.rerun()

    st.caption("Текущий файл tasks.txt (предпросмотр):")
    st.text_area("tasks_preview", value="\n".join(read_tasks_txt()) if read_tasks_txt() else "(пусто)", height=250, label_visibility="collapsed", disabled=True)

# ─── MAIN APP (ГЕНЕРАЦИЯ + QC ВМЕСТЕ) ──────────────────────────────────────────
if not tasks: st.stop()
task = tasks[cur]

st.markdown("### ⚡ Массовые действия (без пауз)")
c1, c2 = st.columns(2)
with c1:
    if st.button("▶️ Заполнить ПУСТЫЕ выбранные (Автоматом)", type="primary", use_container_width=True):
        st.session_state.auto_run_selected = True
        st.rerun()
with c2:
    if st.button("⚠️ ПЕРЕЗАПИСАТЬ выбранные (Заново)", use_container_width=True):
        st.session_state.show_confirm_overwrite = True
        st.rerun()

if st.session_state.get("show_confirm_overwrite", False):
    st.warning(f"Начнется НЕПРЕРЫВНЫЙ авто-прогон для {len(st.session_state.selected_tasks)} папок. Старые данные будут перезаписаны.")
    cw1, cw2 = st.columns(2)
    if cw1.button("🔴 ДА, НАЧАТЬ МАССОВУЮ ПЕРЕЗАПИСЬ", type="primary"):
        st.session_state.show_confirm_overwrite = False
        st.session_state.done_set = st.session_state.done_set - st.session_state.selected_tasks
        
        # Используем JSON для логирования очереди вместо физического удаления Excel-файла
        for t_idx in st.session_state.selected_tasks:
            st.session_state.qc_state[t_idx] = {"TITLES": "redo", "DESCS": "redo", "CTAS": "redo", "TAGS": "wait"}
            save_qc_state(tasks[t_idx]["path"], st.session_state.qc_state[t_idx])

        st.session_state.auto_run_selected = True
        st.rerun()
    if cw2.button("Отмена"):
        st.session_state.show_confirm_overwrite = False
        st.rerun()

st.divider()

def qc_radio(label, current_val, key):
    opts = {"wait": "⏳ Ждет проверки", "ok": "🟢 ОК", "redo": "🔴 Переделать"}
    return st.radio(label, options=list(opts.keys()), format_func=lambda x: opts[x],
                    index=list(opts.keys()).index(current_val), key=key, horizontal=True)

def set_qc(t, d, c, tag=None):
    if tag is None: tag = st.session_state.qc_state[cur].get("TAGS", "wait")
    st.session_state.qc_state[cur] = {"TITLES": t, "DESCS": d, "CTAS": c, "TAGS": tag}
    save_qc_state(task["path"], st.session_state.qc_state[cur])

# ─── ВЕРХНИЙ БЛОК: НАСТРОЙКИ СЛЕВА, ТЕГИ СПРАВА ───────────────────────────────
top_left, top_right = st.columns([6, 4])

with top_left:
    json_file_path = os.path.join(base_dir, "all_products.json")
    if not os.path.exists(json_file_path): json_file_path = os.path.join(task['path'], "all_products.json")
    json_text = f"📎 JSON: {os.path.basename(json_file_path)}" if os.path.exists(json_file_path) else "⚠️ JSON не найден!"

    st.markdown(f"""
    <div class="card card-accent">
      <div class="niche-badge">📁 {task['niche']} ➜ {task['strat']}</div>
      <div style="font-size: 11px; color: #34d399;">📊 {task['xlsx']}</div>
      <div style="font-size: 11px; color: #a78bfa; margin-top: 3px;">{json_text}</div>
    </div>
    """, unsafe_allow_html=True)

    st.session_state.niche_mode = st.checkbox(
        "🧭 Бить в ОБЩУЮ НИШУ (авто-определение темы перед генерацией)",
        value=st.session_state.niche_mode,
        help="Включи, если обычная генерация промахивается мимо части фото из выдачи."
    )
    st.session_state.use_json = st.checkbox(
        "📎 Прикреплять all_products.json к промпту",
        value=st.session_state.use_json
    )

    st.markdown("<div style='height: 10px;'></div>", unsafe_allow_html=True)
    new_tag_val = qc_radio("🏷️ Оценка тегов (Boards_Tags.txt):", st.session_state.qc_state[cur].get("TAGS", "wait"), f"qc_tag_{cur}")
    if new_tag_val != st.session_state.qc_state[cur].get("TAGS", "wait"):
        st.session_state.qc_state[cur]["TAGS"] = new_tag_val
        save_qc_state(task["path"], st.session_state.qc_state[cur])
        st.rerun()
        
    st.markdown("<div style='height: 10px;'></div>", unsafe_allow_html=True)

    with st.expander("🚫 Исключить ниши (по папкам)", expanded=False):
        if st.session_state.niche_exclusion_rules:
            for ridx, rule in enumerate(st.session_state.niche_exclusion_rules):
                rc1, rc2 = st.columns([10, 1])
                with rc1:
                    scope_label = "🌍 ВСЕ папки" if rule["scope"] == "all" else ("📁 " + "; ".join([f"{tasks[i]['niche']} ➜ {tasks[i]['strat']}" for i in rule["scope"]]))
                    st.markdown(f"**{scope_label}**  \n🚫 {rule['text']}")
                with rc2:
                    if st.button("✕", key=f"del_rule_{ridx}"):
                        st.session_state.niche_exclusion_rules.pop(ridx)
                        st.rerun()
            st.divider()

        st.caption("Добавить новое исключение:")
        new_scope_all = st.checkbox("Применить ко ВСЕМ папкам", key="new_rule_all")
        new_scope_tasks = [] if new_scope_all else st.multiselect("Или выбери конкретные папки", options=list(range(len(tasks))), format_func=lambda i: f"{tasks[i]['niche']} ➜ {tasks[i]['strat']}", key="new_rule_tasks")
        new_rule_text = st.text_input("Ниши через запятую", placeholder="Например: Nursery, Baby Shower", key="new_rule_text")
        if st.button("➕ Добавить исключение", key="add_rule_btn"):
            if new_rule_text.strip() and (new_scope_all or new_scope_tasks):
                st.session_state.niche_exclusion_rules.append({"scope": "all" if new_scope_all else new_scope_tasks, "text": new_rule_text.strip()})
                st.rerun()

    with st.expander("💡 Подсказка по нише (по папкам)", expanded=False):
        if st.session_state.niche_hint_rules:
            for ridx, rule in enumerate(st.session_state.niche_hint_rules):
                rc1, rc2 = st.columns([10, 1])
                with rc1:
                    scope_label = "🌍 ВСЕ папки" if rule["scope"] == "all" else ("📁 " + "; ".join([f"{tasks[i]['niche']} ➜ {tasks[i]['strat']}" for i in rule["scope"]]))
                    st.markdown(f"**{scope_label}**  \n💡 {rule['text']}")
                with rc2:
                    if st.button("✕", key=f"del_hint_rule_{ridx}"):
                        st.session_state.niche_hint_rules.pop(ridx)
                        st.rerun()
            st.divider()

        st.caption("Добавить новую подсказку:")
        new_hint_all = st.checkbox("Применить ко ВСЕМ папкам", key="new_hint_all")
        new_hint_tasks = [] if new_hint_all else st.multiselect("Или выбери конкретные папки", options=list(range(len(tasks))), format_func=lambda i: f"{tasks[i]['niche']} ➜ {tasks[i]['strat']}", key="new_hint_tasks")
        new_hint_text = st.text_input("В какую нишу склонять определение", placeholder="Например: Teddy Bear", key="new_hint_text")
        if st.button("➕ Добавить подсказку", key="add_hint_rule_btn"):
            if new_hint_text.strip() and (new_hint_all or new_hint_tasks):
                st.session_state.niche_hint_rules.append({"scope": "all" if new_hint_all else new_hint_tasks, "text": new_hint_text.strip()})
                st.rerun()

    # Отображаем ошибку AI, если она была при последней попытке
    if "_last_ai_error" in st.session_state and st.session_state["_last_ai_error"]:
        st.error(st.session_state["_last_ai_error"])

    def try_auto_save(cur_task_idx):
        t_lines = clean_lines(st.session_state.get("ta_titles", ""))
        d_lines = clean_lines(st.session_state.get("ta_descs", ""))
        c_lines = clean_lines(st.session_state.get("ta_ctas", ""))
        if t_lines and d_lines and c_lines:
            write_xlsx(st.session_state.tasks[cur_task_idx]["xlsx"], t_lines, d_lines, c_lines)

    mc1, mc2 = st.columns(2)
    with mc1:
        if st.button("🤖 Сгенерировать ВСЁ для этой папки", use_container_width=True):
            with st.spinner("Работаем с AI..."):
                if run_ai_for_current_task(cur, "ALL"):
                    try_auto_save(cur)
                    st.rerun()
    with mc2:
        with st.popover("🔄 Точечная регенерация (исправить)"):
            if st.button("Переделать ЗАГОЛОВКИ"):
                if run_ai_for_current_task(cur, "TITLES"):
                    try_auto_save(cur)
                    st.rerun()
            if st.button("Переделать ОПИСАНИЯ"):
                if run_ai_for_current_task(cur, "DESCS"):
                    try_auto_save(cur)
                    st.rerun()
            if st.button("Переделать CTA"):
                if run_ai_for_current_task(cur, "CTAS"):
                    try_auto_save(cur)
                    st.rerun()

def render_tags_copier(tags_content, height=350):
    lines = [line.strip() for line in tags_content.splitlines() if line.strip()]
    
    html_str = """
    <!DOCTYPE html>
    <html>
    <head>
    <style>
    body { 
        margin: 0; 
        padding: 0 5px;
        font-family: 'JetBrains Mono', monospace, sans-serif; 
        background-color: transparent;
    }
    .tag-item {
        display: flex; 
        justify-content: space-between; 
        align-items: center;
        padding: 8px 12px; 
        margin-bottom: 6px;
        background: rgba(139, 92, 246, 0.15);
        border-radius: 6px; 
        border: 1px solid rgba(139, 92, 246, 0.3);
        color: var(--dynamic-text-color, #ffffff); /* Цвет берется из родителя */
        font-size: 13px; 
        cursor: pointer; 
        transition: all 0.2s ease;
    }
    .tag-item:hover { 
        background: rgba(139, 92, 246, 0.3); 
        border-color: rgba(139, 92, 246, 0.6);
    }
    .tag-item:active { 
        background: rgba(139, 92, 246, 0.5); 
    }
    .copy-icon { 
        font-size: 14px; 
        opacity: 0.7; 
    }
    </style>
    <script>
    document.addEventListener("DOMContentLoaded", () => {
        try {
            // Подтягиваем реальный цвет текста из интерфейса Streamlit (чтобы работало при любой теме)
            const parentColor = window.parent.getComputedStyle(window.parent.document.body).color;
            document.documentElement.style.setProperty('--dynamic-text-color', parentColor);
        } catch(e) {
            console.log("CORS prevented reading parent color, fallback to white");
        }
    });

    function copyTag(el, text) {
        const textArea = document.createElement("textarea");
        textArea.value = text;
        textArea.style.position = "fixed";
        document.body.appendChild(textArea);
        textArea.focus();
        textArea.select();
        try {
            document.execCommand('copy');
            let icon = el.querySelector('.copy-icon');
            icon.innerText = '✅';
            setTimeout(() => { icon.innerText = '📋'; }, 1200);
        } catch (err) {
            console.error('Fallback: Oops, unable to copy', err);
        }
        document.body.removeChild(textArea);
    }
    </script>
    </head>
    <body>
    """
    for tag in lines:
        safe_tag = tag.replace("'", "\\'").replace('"', '&quot;')
        display_tag = tag.replace('<', '&lt;').replace('>', '&gt;')
        html_str += f"""<div class="tag-item" onclick="copyTag(this, '{safe_tag}')"><span>{display_tag}</span><span class="copy-icon">📋</span></div>\n"""        
    html_str += "</body></html>"
    components.html(html_str, height=height, scrolling=True)

with top_right:
    tags_path = os.path.join(task['path'], "For Main", "Boards_Tags.txt")
    if os.path.exists(tags_path):
        tags_content = read_txt(tags_path)
        
        # Формируем полный красивый путь (обрезаем всё до Managers Pinterest)
        if "Managers Pinterest/" in tags_path:
            display_path = tags_path.split("Managers Pinterest/")[-1]
        else:
            display_path = os.path.relpath(tags_path, base_dir)
            
        st.markdown(f"**🏷️ Теги (Boards_Tags.txt) — {len(tags_content.splitlines())} шт.**")
        st.code(display_path, language="text")
        render_tags_copier(tags_content, height=350)
    else:
        st.warning("⚠️ Файл Boards_Tags.txt не найден в этой папке.")

st.divider()

# --- ПОЛЯ И КОНТРОЛЬ КАЧЕСТВА ---
col1, col2, col3 = st.columns(3)

with col1:
    badge_t = " 🆕 ТОЛЬКО ЧТО ПЕРЕДЕЛАНО" if "TITLES" in st.session_state.just_regenerated.get(cur, set()) else ""
    raw_t = st.text_area(f"📝 ЗАГОЛОВКИ{badge_t}", value=st.session_state.get("ta_titles", ""), height=340, key=f"ta_titles_w_{st.session_state.render_keys['ta_titles']}")
    if raw_t != st.session_state.get("ta_titles", ""): st.session_state["ta_titles"] = raw_t
    titles = clean_lines(raw_t)
    st.caption(f"Строк: {len(titles)}")
    new_t = qc_radio("Оценка:", st.session_state.qc_state[cur]["TITLES"], f"qct_{cur}")

with col2:
    badge_d = " 🆕 ТОЛЬКО ЧТО ПЕРЕДЕЛАНО" if "DESCS" in st.session_state.just_regenerated.get(cur, set()) else ""
    raw_d = st.text_area(f"📖 ОПИСАНИЯ{badge_d}", value=st.session_state.get("ta_descs", ""), height=340, key=f"ta_descs_w_{st.session_state.render_keys['ta_descs']}")
    if raw_d != st.session_state.get("ta_descs", ""): st.session_state["ta_descs"] = raw_d
    descs = clean_lines(raw_d)
    st.caption(f"Строк: {len(descs)}")
    new_d = qc_radio("Оценка:", st.session_state.qc_state[cur]["DESCS"], f"qcd_{cur}")

with col3:
    badge_c = " 🆕 ТОЛЬКО ЧТО ПЕРЕДЕЛАНО" if "CTAS" in st.session_state.just_regenerated.get(cur, set()) else ""
    raw_c = st.text_area(f"🎯 CTA{badge_c}", value=st.session_state.get("ta_ctas", ""), height=340, key=f"ta_ctas_w_{st.session_state.render_keys['ta_ctas']}")
    if raw_c != st.session_state.get("ta_ctas", ""): st.session_state["ta_ctas"] = raw_c
    ctas = clean_lines(raw_c)
    st.caption(f"Строк: {len(ctas)}")
    new_c = qc_radio("Оценка:", st.session_state.qc_state[cur]["CTAS"], f"qcc_{cur}")

if new_t != st.session_state.qc_state[cur]["TITLES"] or new_d != st.session_state.qc_state[cur]["DESCS"] or new_c != \
        st.session_state.qc_state[cur]["CTAS"]:
    if new_t != st.session_state.qc_state[cur]["TITLES"]: st.session_state.just_regenerated.get(cur, set()).discard("TITLES")
    if new_d != st.session_state.qc_state[cur]["DESCS"]: st.session_state.just_regenerated.get(cur, set()).discard("DESCS")
    if new_c != st.session_state.qc_state[cur]["CTAS"]: st.session_state.just_regenerated.get(cur, set()).discard("CTAS")
    set_qc(new_t, new_d, new_c, st.session_state.qc_state[cur].get("TAGS", "wait"))
    st.rerun()

st.divider()

# ─── ГЛОБАЛЬНЫЕ КНОПКИ QC И СОХРАНЕНИЯ ────────────────────────────────────────
qc_b1, qc_b2 = st.columns(2)
if qc_b1.button("✅ ПОМЕТИТЬ ТЕКСТЫ ИДЕАЛЬНЫМИ (Без Тегов)", use_container_width=True):
    set_qc("ok", "ok", "ok")
    st.rerun()
if qc_b2.button("🔴 ОТПРАВИТЬ ТЕКСТЫ НА ПЕРЕДЕЛКУ", use_container_width=True):
    set_qc("redo", "redo", "redo")
    st.rerun()

DUP_WARNING_THRESHOLD = 10 

def _dup_warning(label, lines):
    uniq = len(set(lines))
    dup_count = len(lines) - uniq
    if dup_count > DUP_WARNING_THRESHOLD:
        st.warning(f"⚠️ В блоке {label} есть {dup_count} дублирующихся строк(и).")

_dup_warning("ЗАГОЛОВКИ", titles)
_dup_warning("ОПИСАНИЯ", descs)
_dup_warning("CTA", ctas)

write_ready = bool(titles and descs and ctas and task["xlsx"])
if st.button("💾 СОХРАНИТЬ ТЕКУЩУЮ СТРАНИЦУ В EXCEL", type="primary", use_container_width=True,
             disabled=not write_ready):
    write_xlsx(task["xlsx"], titles, descs, ctas)
    st.session_state.done_set.add(cur)
    st.success("✅ Сохранено в Excel!")

st.divider()

# ─── БЛОК АВТОМАТИЧЕСКОГО ИСПРАВЛЕНИЯ БРАКА ─────────────────────────────────
total_redos = sum(1 for q in st.session_state.qc_state.values() for v in q.values() if v == "redo")
if total_redos > 0:
    st.markdown(f"**Блоков, отправленных на переделку (по всем папкам):** <span style='color:#f87171;font-weight:bold;'>{total_redos}</span>", unsafe_allow_html=True)
    if st.button("🚀 ЗАВЕРШИТЬ КОНТРОЛЬ И ИСПРАВИТЬ БРАК (Автоматом через ИИ)", type="primary", use_container_width=True):
        st.session_state.show_confirm_qc = True
        st.rerun()

    if st.session_state.get("show_confirm_qc", False):
        st.warning(f"Начнется непрерывный авто-прогон для {total_redos} бракованных блоков.")
        c_qc1, c_qc2 = st.columns(2)
        if c_qc1.button("✅ ДА, ИСПРАВИТЬ ВСЁ"):
            st.session_state.show_confirm_qc = False
            st.session_state.qc_run_redos = True
            st.rerun()
        if c_qc2.button("❌ Отмена"):
            st.session_state.show_confirm_qc = False
            st.rerun()